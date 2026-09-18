#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o Produto 01 — Diagnóstico da Governança Contratual.

Fonte única de conteúdo: BASE_CONTEXTO_CONSULTORIA_TIC.md (documento base
do projeto, versão 1.1). Nenhuma informação é acrescentada por inferência:
o que o documento base registra como pendente permanece pendente aqui.
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor

from metricas_fonte import (MARGEM_CELULA_CM, MARGEM_CELULA_TWIPS,
                            largura_minima_cm)

# --------------------------------------------------------------------------
# Identidade do documento
# --------------------------------------------------------------------------

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SAIDA_DOCX = os.path.join(
    RAIZ, "Produto_01_Diagnostico_Governanca_Contratual.docx"
)

FONTE = "Arial"
AZUL = RGBColor(0x1F, 0x35, 0x5E)       # títulos e filetes
GRAFITE = RGBColor(0x33, 0x33, 0x33)    # corpo de texto
CINZA = RGBColor(0x59, 0x59, 0x59)      # textos auxiliares
FILL_CAB = "1F355E"                     # fundo do cabeçalho das tabelas
FILL_SUB = "DCE3EE"                     # fundo de linhas de subtotal/destaque
FILL_ZEBRA = "F2F5F9"                   # fundo alternado

TITULO = "Diagnóstico da Governança Contratual"
PRODUTO = "Produto 01"
CONSULTORA = "Alana Leitielly da Silva"
CONTRATO = "ED01318/2026"
DATA_ELAB = "18 de setembro de 2026"
SEI_DATAPREV = "23000.039481/2024-24"
SEI_SERPRO = "23000.005397/2025-98"
PEND = "[PENDENTE]"
VALID = "[A VALIDAR]"


# --------------------------------------------------------------------------
# Utilitários de baixo nível
# --------------------------------------------------------------------------

def _el(tag, **attrs):
    e = OxmlElement(tag)
    for k, v in attrs.items():
        e.set(qn("w:" + k), v)
    return e


def campo(paragrafo, instrucao, cache=""):
    """Insere um campo do Word (PAGE, NUMPAGES, TOC) com resultado em cache."""
    r = paragrafo.add_run()
    r._r.append(_el("w:fldChar", fldCharType="begin"))
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instrucao
    r._r.append(it)
    r._r.append(_el("w:fldChar", fldCharType="separate"))
    rc = paragrafo.add_run(cache)
    rc._r.append(_el("w:fldChar", fldCharType="end"))
    return r, rc


def sombrear(celula, cor):
    celula._tc.get_or_add_tcPr().append(
        _el("w:shd", val="clear", color="auto", fill=cor)
    )


def repetir_cabecalho(linha):
    linha._tr.get_or_add_trPr().append(_el("w:tblHeader", val="true"))


def nao_dividir(linha):
    """Impede que a linha seja partida entre duas páginas."""
    linha._tr.get_or_add_trPr().append(_el("w:cantSplit", val="true"))


def bordas(tabela, cor="B7C0D0", tamanho="4"):
    tblPr = tabela._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b.append(_el("w:" + lado, val="single", sz=tamanho,
                     space="0", color=cor))
    tblPr.append(b)


# Ordem exigida pelo esquema OOXML para os filhos de <w:tblPr>.
ORDEM_TBLPR = [
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
    "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
    "tblCaption", "tblDescription",
]


def ordenar_tblPr(tblPr):
    """Reordena os filhos de <w:tblPr> e remove duplicatas, para que o
    arquivo seja válido também no Word, que é estrito quanto à ordem."""
    posicao = {qn("w:" + nome): i for i, nome in enumerate(ORDEM_TBLPR)}
    filhos, vistos = [], set()
    for filho in list(tblPr):
        tblPr.remove(filho)
        if filho.tag in posicao and filho.tag in vistos:
            continue
        vistos.add(filho.tag)
        filhos.append(filho)
    for filho in sorted(filhos, key=lambda e: posicao.get(e.tag, 99)):
        tblPr.append(filho)


def definir_margem_celulas(tabela, twips=MARGEM_CELULA_TWIPS):
    margens = _el("w:tblCellMar")
    for lado in ("top", "left", "bottom", "right"):
        valor = "20" if lado in ("top", "bottom") else str(twips)
        margens.append(_el("w:" + lado, w=valor, type="dxa"))
    tabela._tbl.tblPr.append(margens)


def ajustar_proporcoes(proporcoes, minimos_cm, largura_total_cm):
    """Corrige as proporções pedidas para que nenhuma coluna fique abaixo
    da largura mínima necessária, retirando o excedente das colunas que
    têm folga, na medida da folga de cada uma."""
    soma = float(sum(proporcoes))
    larguras = [largura_total_cm * peso / soma for peso in proporcoes]
    if sum(minimos_cm) > largura_total_cm:
        print("  aviso: tabela sem largura suficiente (%.1f cm necessários, "
              "%.1f cm disponíveis)" % (sum(minimos_cm), largura_total_cm))
        return minimos_cm

    for _ in range(40):
        deficit = sum(max(0.0, m - l) for m, l in zip(minimos_cm, larguras))
        if deficit < 0.001:
            break
        folgas = [max(0.0, l - m) for l, m in zip(larguras, minimos_cm)]
        total_folga = sum(folgas)
        if total_folga < 0.001:
            break
        ajustadas = []
        for larg, minimo, folga in zip(larguras, minimos_cm, folgas):
            if larg < minimo:
                ajustadas.append(minimo)
            else:
                retirada = deficit * folga / total_folga
                ajustadas.append(max(minimo, larg - retirada))
        larguras = ajustadas
    return larguras


def largura_fixa(tabela, proporcoes, largura_total_cm=16.0,
                 absolutas=False):
    """Fixa a largura das colunas segundo proporções relativas, tanto na
    grade da tabela quanto em cada célula."""
    tabela.autofit = False
    if absolutas:
        larguras = [Cm(v) for v in proporcoes]
        largura_total_cm = sum(proporcoes)
    else:
        soma = float(sum(proporcoes))
        larguras = [Cm(largura_total_cm * peso / soma)
                    for peso in proporcoes]

    grade = tabela._tbl.find(qn("w:tblGrid"))
    if grade is not None:
        for coluna_grade, larg in zip(grade.findall(qn("w:gridCol")),
                                      larguras):
            coluna_grade.set(qn("w:w"), str(Emu(larg).twips))

    for coluna, larg in zip(tabela.columns, larguras):
        for celula in coluna.cells:
            celula.width = larg

    tblW = tabela._tbl.tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = _el("w:tblW")
        tabela._tbl.tblPr.append(tblW)
    tblW.set(qn("w:w"), str(Emu(Cm(largura_total_cm)).twips))
    tblW.set(qn("w:type"), "dxa")


def espacamento_celula(tabela, antes=2, depois=2):
    for linha in tabela.rows:
        for celula in linha.cells:
            for p in celula.paragraphs:
                p.paragraph_format.space_before = Pt(antes)
                p.paragraph_format.space_after = Pt(depois)


# --------------------------------------------------------------------------
# Estilos
# --------------------------------------------------------------------------

def configurar_estilos(doc):
    estilos = doc.styles

    normal = estilos["Normal"]
    normal.font.name = FONTE
    normal.font.size = Pt(11)
    normal.font.color.rgb = GRAFITE
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for atr in ("ascii", "hAnsi", "cs", "eastAsia"):
        rfonts.set(qn("w:" + atr), FONTE)
    rpr.append(_el("w:lang", val="pt-BR", eastAsia="pt-BR", bidi="ar-SA"))
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_after = Pt(8)
    pf.line_spacing = 1.3

    cfg = {
        "Heading 1": (15, True, 20, 10, True),
        "Heading 2": (12.5, True, 16, 6, False),
        "Heading 3": (11.5, True, 12, 4, False),
    }
    for nome, (tam, negrito, antes, depois, quebra) in cfg.items():
        est = estilos[nome]
        est.font.name = FONTE
        est.font.size = Pt(tam)
        est.font.bold = negrito
        est.font.color.rgb = AZUL
        est.font.italic = False
        r = est.element.get_or_add_rPr()
        rf = r.get_or_add_rFonts()
        for atr in ("ascii", "hAnsi", "cs", "eastAsia"):
            rf.set(qn("w:" + atr), FONTE)
        p = est.paragraph_format
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.space_before = Pt(antes)
        p.space_after = Pt(depois)
        p.line_spacing = 1.1
        p.keep_with_next = True
        p.page_break_before = False

    def novo(nome, base="Normal"):
        if nome in [s.name for s in estilos]:
            return estilos[nome]
        return estilos.add_style(nome, 1)  # WD_STYLE_TYPE.PARAGRAPH

    st = novo("CorpoTabela")
    st.base_style = estilos["Normal"]
    st.font.size = Pt(9)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.05

    st = novo("NotaTecnica")
    st.base_style = estilos["Normal"]
    st.font.size = Pt(9.5)
    st.font.color.rgb = CINZA
    st.paragraph_format.left_indent = Cm(0.6)
    st.paragraph_format.space_before = Pt(4)
    st.paragraph_format.space_after = Pt(10)
    st.paragraph_format.line_spacing = 1.2

    st = novo("Legenda")
    st.base_style = estilos["Normal"]
    st.font.size = Pt(9)
    st.font.bold = True
    st.font.color.rgb = AZUL
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_before = Pt(10)
    st.paragraph_format.space_after = Pt(3)
    st.paragraph_format.keep_with_next = True

    st = novo("Ficha")
    st.base_style = estilos["Normal"]
    st.font.name = "Courier New"
    st.font.size = Pt(9)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_after = Pt(0)
    st.paragraph_format.line_spacing = 1.0
    st.paragraph_format.left_indent = Cm(0.4)


# --------------------------------------------------------------------------
# Blocos de construção
# --------------------------------------------------------------------------

def h(doc, texto, nivel=1, nova_pagina=False):
    p = doc.add_heading(texto, level=nivel)
    if nova_pagina:
        p.paragraph_format.page_break_before = True
    return p


def par(doc, texto, estilo=None, alinhamento=None, antes=None, depois=None):
    p = doc.add_paragraph(texto, style=estilo)
    if alinhamento is not None:
        p.alignment = alinhamento
    if antes is not None:
        p.paragraph_format.space_before = Pt(antes)
    if depois is not None:
        p.paragraph_format.space_after = Pt(depois)
    return p


def nota(doc, rotulo, texto):
    """Nota de seção, introduzida por 'Nota.' ou 'Ressalva.' em destaque."""
    p = doc.add_paragraph(style="NotaTecnica")
    r = p.add_run(rotulo + " ")
    r.bold = True
    r.font.color.rgb = AZUL
    p.add_run(texto)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    borda = _el("w:pBdr")
    borda.append(_el("w:left", val="single", sz="12", space="8", color="1F355E"))
    p._p.get_or_add_pPr().append(borda)
    return p


_CONTADOR_QUADRO = [0]


def legenda(doc, texto):
    """Rotula o quadro seguinte, numerando-o na ordem do documento."""
    _CONTADOR_QUADRO[0] += 1
    return par(doc, "Quadro %d — %s" % (_CONTADOR_QUADRO[0], texto),
               estilo="Legenda")


def larguras_minimas(cabecalho, linhas, fonte):
    """Largura mínima de cada coluna para que nenhuma palavra seja
    partida: o cabeçalho é negrito, e as marcações de situação também."""
    minimos = [largura_minima_cm(texto, fonte, negrito=True)
               for texto in cabecalho]
    for linha in linhas:
        for i, texto in enumerate(linha):
            if i >= len(minimos):
                continue
            texto = str(texto)
            negrito = texto.strip() in (PEND, VALID, "[VEDADO]")
            minimos[i] = max(minimos[i],
                             largura_minima_cm(texto, fonte, negrito))
    return [m + MARGEM_CELULA_CM + 0.04 for m in minimos]


def tabela(doc, cabecalho, linhas, proporcoes=None, fonte=9,
           largura=16.0, zebra=True, alinhamentos=None):
    t = doc.add_table(rows=1, cols=len(cabecalho))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    bordas(t)

    cab = t.rows[0]
    repetir_cabecalho(cab)
    nao_dividir(cab)
    for i, texto in enumerate(cabecalho):
        celula = cab.cells[i]
        celula.text = ""
        p = celula.paragraphs[0]
        p.style = doc.styles["CorpoTabela"]
        r = p.add_run(texto)
        r.bold = True
        r.font.size = Pt(fonte)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        sombrear(celula, FILL_CAB)

    for n, linha in enumerate(linhas):
        nova = t.add_row()
        nao_dividir(nova)
        celulas = nova.cells
        for i, texto in enumerate(linha):
            celula = celulas[i]
            celula.text = ""
            p = celula.paragraphs[0]
            p.style = doc.styles["CorpoTabela"]
            if alinhamentos and alinhamentos[i] == "c":
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for j, trecho in enumerate(str(texto).split("\n")):
                if j:
                    p = celula.add_paragraph(style="CorpoTabela")
                r = p.add_run(trecho)
                r.font.size = Pt(fonte)
                if trecho.strip() in (PEND, VALID, "[VEDADO]"):
                    r.font.color.rgb = RGBColor(0x8A, 0x3A, 0x1F)
                    r.bold = True
            if zebra and n % 2 == 1:
                sombrear(celula, FILL_ZEBRA)

    if proporcoes:
        minimos = larguras_minimas(cabecalho, linhas, fonte)
        larguras = ajustar_proporcoes(proporcoes, minimos, largura)
        largura_fixa(t, larguras, absolutas=True)
    definir_margem_celulas(t)
    ordenar_tblPr(t._tbl.tblPr)
    espacamento_celula(t)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def ficha(doc, linhas):
    """Bloco de formulário em fonte monoespaçada, com moldura."""
    t = doc.add_table(rows=1, cols=1)
    bordas(t, cor="8A96AC")
    celula = t.rows[0].cells[0]
    sombrear(celula, "F7F8FA")
    celula.text = ""
    primeiro = True
    for texto in linhas:
        p = celula.paragraphs[0] if primeiro else celula.add_paragraph()
        primeiro = False
        p.style = doc.styles["Ficha"]
        r = p.add_run(texto)
        r.font.name = "Courier New"
        r.font.size = Pt(9)
        if texto and not texto.startswith(" ") and texto.isupper():
            r.bold = True
    largura_fixa(t, [1], 16.0)
    ordenar_tblPr(t._tbl.tblPr)
    espacamento_celula(t, antes=1, depois=1)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


# --------------------------------------------------------------------------
# Página, cabeçalho e rodapé
# --------------------------------------------------------------------------

def configurar_pagina(doc):
    s = doc.sections[0]
    s.page_width = Cm(21.0)
    s.page_height = Cm(29.7)
    s.left_margin = Cm(2.5)
    s.right_margin = Cm(2.5)
    s.top_margin = Cm(2.6)
    s.bottom_margin = Cm(2.2)
    s.header_distance = Cm(1.2)
    s.footer_distance = Cm(1.0)
    s.different_first_page_header_footer = True
    return s


def nova_secao(doc, paisagem):
    """Abre uma seção com orientação própria, preservando cabeçalho e
    rodapé. Usada para acomodar as tabelas largas em paisagem."""
    s = doc.add_section(WD_SECTION.NEW_PAGE)
    if paisagem:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin, s.bottom_margin = Cm(2.0), Cm(1.8)
    else:
        s.orientation = WD_ORIENT.PORTRAIT
        s.page_width, s.page_height = Cm(21.0), Cm(29.7)
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin, s.bottom_margin = Cm(2.6), Cm(2.2)
    s.header_distance, s.footer_distance = Cm(1.2), Cm(1.0)
    s.different_first_page_header_footer = False
    s.header.is_linked_to_previous = False
    s.footer.is_linked_to_previous = False
    montar_cabecalho(s)
    montar_rodape(s, Emu(s.page_width - s.left_margin
                             - s.right_margin).twips)
    return s


def montar_cabecalho(secao):
    cab = secao.header
    p = cab.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run("%s — %s" % (PRODUTO, TITULO))
    r.bold = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = AZUL
    r.font.name = FONTE

    p2 = cab.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(
        "Consultoria Técnica em Governança de Contratos de TIC   ·   "
        "UNESCO   ·   Ministério da Educação"
    )
    r2.font.size = Pt(8)
    r2.font.color.rgb = CINZA
    r2.font.name = FONTE
    borda = _el("w:pBdr")
    borda.append(_el("w:bottom", val="single", sz="6", space="4",
                     color="1F355E"))
    p2._p.get_or_add_pPr().append(borda)


def montar_rodape(secao, tabulacao=9071):
    rod = secao.footer
    p = rod.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    borda = _el("w:pBdr")
    borda.append(_el("w:top", val="single", sz="6", space="6", color="B7C0D0"))
    p._p.get_or_add_pPr().append(borda)

    tabs = _el("w:tabs")
    tabs.append(_el("w:tab", val="right", pos=str(tabulacao)))
    p._p.get_or_add_pPr().append(tabs)

    r = p.add_run(
        "%s   ·   Contrato UNESCO %s   ·   Documento técnico\t"
        % (CONSULTORA, CONTRATO)
    )
    r.font.size = Pt(8)
    r.font.color.rgb = CINZA
    r.font.name = FONTE

    r = p.add_run("Página ")
    r.font.size = Pt(8)
    r.font.color.rgb = CINZA
    r.font.name = FONTE
    campo(p, "PAGE", "2")
    r = p.add_run(" de ")
    r.font.size = Pt(8)
    r.font.color.rgb = CINZA
    r.font.name = FONTE
    campo(p, "NUMPAGES", "40")
    for run in p.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = CINZA
        run.font.name = FONTE


# --------------------------------------------------------------------------
# Elementos pré-textuais
# --------------------------------------------------------------------------

def capa(doc):
    def linha(texto, tam, negrito=False, cor=GRAFITE, antes=0, depois=6,
              espacamento=None, maiuscula=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(antes)
        p.paragraph_format.space_after = Pt(depois)
        p.paragraph_format.line_spacing = espacamento or 1.15
        r = p.add_run(texto.upper() if maiuscula else texto)
        r.bold = negrito
        r.font.size = Pt(tam)
        r.font.color.rgb = cor
        r.font.name = FONTE
        return p

    def filete(largura_pt="12", cor="1F355E", antes=10, depois=10):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(antes)
        p.paragraph_format.space_after = Pt(depois)
        b = _el("w:pBdr")
        b.append(_el("w:bottom", val="single", sz=largura_pt, space="1",
                     color=cor))
        p._p.get_or_add_pPr().append(b)
        return p

    linha("Ministério da Educação", 12, True, AZUL, antes=18, depois=2,
          maiuscula=True)
    linha(
        "Organização das Nações Unidas para a Educação, a Ciência e a Cultura",
        10, False, CINZA, depois=2,
    )

    filete(antes=30, depois=40)

    linha("Produto 01", 13, True, CINZA, depois=10, maiuscula=True)
    linha("Diagnóstico da Governança Contratual", 24, True, AZUL, depois=14,
          espacamento=1.05)
    linha(
        "Contratos estratégicos de Tecnologia da Informação e Comunicação "
        "celebrados pelo Ministério da Educação com a DATAPREV e com o SERPRO",
        12, False, GRAFITE, depois=6, espacamento=1.25,
    )

    filete(largura_pt="6", cor="B7C0D0", antes=18, depois=18)

    linha("Processos analisados", 9.5, True, CINZA, depois=3, maiuscula=True)
    linha("DATAPREV — Processo SEI nº %s" % SEI_DATAPREV, 11, False, GRAFITE,
          depois=2)
    linha("SERPRO — Processo SEI nº %s" % SEI_SERPRO, 11, False, GRAFITE,
          depois=0)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.LINE)

    linha("Consultoria Técnica em Governança de Contratos de TIC", 12, True,
          AZUL, antes=90, depois=4)
    linha("Atividades contratuais 1.1 e 1.2", 10.5, False, CINZA, depois=22)

    linha("Consultora Individual", 9.5, True, CINZA, depois=2,
          maiuscula=True)
    linha(CONSULTORA, 13, True, GRAFITE, depois=4)
    linha("Contrato de Consultoria UNESCO nº %s" % CONTRATO, 10.5, False,
          CINZA, depois=0)

    filete(largura_pt="6", cor="B7C0D0", antes=48, depois=12)
    linha("2026", 12, True, AZUL, depois=0)


def identificacao(doc):
    h(doc, "IDENTIFICAÇÃO DO PRODUTO", 1, nova_pagina=True)
    par(
        doc,
        "O quadro a seguir identifica o produto, vincula-o às atividades "
        "previstas no contrato de consultoria e registra os dados "
        "institucionais da contratação. Os campos assinalados como "
        "pendentes não foram informados até a data de elaboração deste "
        "documento e não são preenchidos por inferência.",
    )
    dados = [
        ("Produto", "Produto 01 — Diagnóstico da Governança Contratual"),
        ("Atividades contratuais",
         "Atividade 1.1 — Definir metodologia de análise da governança "
         "contratual.\nAtividade 1.2 — Elaborar diagnóstico técnico da "
         "governança contratual."),
        ("Consultora individual", CONSULTORA),
        ("Contrato de consultoria", "UNESCO nº %s" % CONTRATO),
        ("Número da SA", "SA-1630/2026"),
        ("Vigência do contrato", "09/09/2026 a 31/12/2026"),
        ("Prazo contratual de entrega do produto", "12/10/2026"),
        ("Supervisão", "Oficial de Projetos do Setor Educação — UNESCO"),
        ("Instituição beneficiária", "Ministério da Educação (MEC)"),
        ("Unidade demandante", PEND),
        ("Projeto de cooperação técnica", PEND),
        ("Objeto da consultoria",
         "Diagnóstico, análise crítica e formulação de metodologia para "
         "avaliação da governança, da conformidade e da eficiência "
         "econômico-financeira dos contratos estratégicos de TIC do "
         "Ministério da Educação, com ênfase nos instrumentos celebrados "
         "com empresas públicas."),
        ("Contratos analisados",
         "DATAPREV — Processo SEI nº %s\nSERPRO — Processo SEI nº %s"
         % (SEI_DATAPREV, SEI_SERPRO)),
        ("Natureza do documento",
         "Documento técnico de consultoria, de uso restrito, elaborado com "
         "base exclusivamente documental."),
        ("Versão", "1.0"),
        ("Data de elaboração", DATA_ELAB),
        ("Local e data de emissão",
         "A completar no ato da assinatura (ver Anexo A, seção A.4)."),
    ]
    tabela(doc, ["Campo", "Informação"], dados, proporcoes=[1, 2.4],
           fonte=9.5, zebra=False)
    nota(
        doc, "Nota.",
        "Em atenção ao artigo VI dos Termos e Condições Gerais do contrato "
        "de consultoria, que veda anunciar, exibir ou apropriar-se do nome, "
        "do logotipo ou do selo oficial da UNESCO, este documento não "
        "reproduz identidade visual institucional. As instituições são "
        "referidas apenas nominalmente, para fins de identificação da "
        "contratação e do objeto analisado.",
    )


def sumario(doc):
    h(doc, "SUMÁRIO", 1, nova_pagina=True)
    par(
        doc,
        "Sumário analítico gerado automaticamente a partir da estrutura de "
        "títulos do documento. Para atualizá-lo no editor de texto, "
        "selecione o sumário e acione o comando de atualização de campos.",
        estilo="NotaTecnica",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    campo(
        p,
        'TOC \\o "1-3" \\h \\z \\u',
        "Sumário a atualizar no editor de texto.",
    )


def siglas(doc):
    h(doc, "LISTA DE SIGLAS E ABREVIATURAS", 1, nova_pagina=True)
    par(
        doc,
        "Relação das siglas empregadas neste documento, com o respectivo "
        "significado.",
    )
    itens = [
        ("COBIT", "Control Objectives for Information and Related "
                  "Technologies"),
        ("DATAPREV", "Empresa pública federal de tecnologia da informação "
                     "contratada no Processo SEI nº %s" % SEI_DATAPREV),
        ("ETP", "Estudo Técnico Preliminar"),
        ("GSI", "Gabinete de Segurança Institucional da Presidência da "
                "República"),
        ("IN", "Instrução Normativa"),
        ("ITIL", "Information Technology Infrastructure Library"),
        ("LGPD", "Lei Geral de Proteção de Dados Pessoais"),
        ("MEC", "Ministério da Educação"),
        ("OS", "Ordem de Serviço"),
        ("PDTIC", "Plano Diretor de Tecnologia da Informação e Comunicação"),
        ("RACI", "Matriz de atribuição de responsabilidades: executa, "
                 "aprova, é consultado e é informado"),
        ("SEI", "Sistema Eletrônico de Informações"),
        ("SERPRO", "Serviço Federal de Processamento de Dados"),
        ("SISP", "Sistema de Administração dos Recursos de Tecnologia da "
                 "Informação"),
        ("TCU", "Tribunal de Contas da União"),
        ("TIC", "Tecnologia da Informação e Comunicação"),
        ("TR", "Termo de Referência"),
        ("UNESCO", "Organização das Nações Unidas para a Educação, a "
                   "Ciência e a Cultura"),
    ]
    tabela(doc, ["Sigla", "Significado"], itens, proporcoes=[1, 4.4],
           fonte=9.5)
    nota(
        doc, "Nota.",
        "As razões sociais completas e atualizadas das empresas contratadas "
        "serão confirmadas a partir dos instrumentos contratuais, quando "
        "disponibilizados, e registradas na versão do produto que "
        "incorporar a análise documental.",
    )


def sumario_executivo(doc):
    h(doc, "SUMÁRIO EXECUTIVO", 1, nova_pagina=True)
    blocos = [
        "Este documento constitui o Produto 01 da consultoria técnica "
        "contratada pela Organização das Nações Unidas para a Educação, a "
        "Ciência e a Cultura, por meio do Contrato nº %s, e tem por objeto o "
        "diagnóstico da governança dos contratos estratégicos de Tecnologia "
        "da Informação e Comunicação celebrados pelo Ministério da Educação "
        "com a DATAPREV, tratada no Processo SEI nº %s, e com o SERPRO, "
        "tratado no Processo SEI nº %s."
        % (CONTRATO, SEI_DATAPREV, SEI_SERPRO),

        "O produto responde a duas atividades contratuais distintas e "
        "sucessivas. A Atividade 1.1 exige a definição da metodologia de "
        "análise da governança contratual, contemplando critérios de "
        "avaliação, dimensões de análise, fontes documentais e parâmetros "
        "normativos aplicáveis. A Atividade 1.2 exige a elaboração do "
        "diagnóstico técnico propriamente dito, com identificação de "
        "lacunas, sobreposições, fragilidades, riscos e oportunidades de "
        "aprimoramento, acompanhada de matriz analítica e de recomendações "
        "estruturadas e mensuráveis.",

        "A Atividade 1.1 encontra-se integralmente desenvolvida nesta "
        "versão. O Capítulo 3 apresenta o instrumento analítico completo e "
        "pronto para aplicação: duas dimensões de análise, a estratégica e "
        "a operacional, decompostas em onze eixos internos; vinte e seis "
        "critérios de avaliação, cada qual com escala própria e regra de "
        "atribuição definida; nove categorias de fontes documentais; o "
        "conjunto de parâmetros normativos aplicáveis, com a situação de "
        "verificação de cada referência; os referenciais de boas práticas e "
        "a regra que disciplina o seu uso; a escala de severidade dos "
        "achados; e os padrões de registro que asseguram a rastreabilidade "
        "de cada afirmação até o documento que a fundamenta.",

        "O achado central desta etapa diz respeito à Atividade 1.2. Até a "
        "data de elaboração deste documento, nenhuma peça dos autos dos dois "
        "processos foi disponibilizada à consultoria. Não foram obtidos os "
        "instrumentos contratuais, os termos de referência, os estudos "
        "técnicos preliminares, os atos de designação de gestores e fiscais, "
        "os registros das instâncias de coordenação, os artefatos de "
        "fiscalização e recebimento nem os mapas de gerenciamento de riscos. "
        "O diagnóstico depende integralmente dessa base documental, e a "
        "metodologia ora consolidada veda de forma expressa a formulação de "
        "achados, de conclusões ou de recomendações sem evidência "
        "documental identificada.",

        "A consequência é direta e deve ser registrada sem atenuação: o "
        "diagnóstico da Atividade 1.2 está estruturado, porém não concluído. "
        "Os capítulos 4 a 7 apresentam os instrumentos de registro já "
        "formatados e prontos para preenchimento — os quadros de verificação "
        "por eixo de análise, as matrizes de lacunas, sobreposições, "
        "fragilidades, riscos e oportunidades, a matriz analítica, as "
        "matrizes de atribuição de responsabilidades e o quadro de "
        "recomendações —, todos com a marcação de situação que identifica "
        "precisamente o que depende de documento ainda não disponibilizado. "
        "Nenhum campo foi preenchido por inferência, por plausibilidade ou "
        "por transposição de conteúdo de produto elaborado para outra "
        "consultoria.",

        "A síntese das recomendações desta etapa é, por consequência, de "
        "natureza procedimental, e está detalhada no Capítulo 8. "
        "Recomenda-se a disponibilização dos documentos relacionados no "
        "Anexo A, observada a ordem de prioridade ali estabelecida, que "
        "concentra na primeira e na segunda faixas as peças capazes de "
        "destravar a maior parte da análise; a resolução das questões "
        "abertas relativas à identificação oficial das referências "
        "normativas indicadas no documento de orientação e à versão dos "
        "referenciais de boas práticas a ser adotada uniformemente nos três "
        "produtos; e a definição do período de execução contratual que "
        "delimitará o universo documental a examinar. Atendidas essas "
        "condições, o diagnóstico poderá ser concluído sem qualquer "
        "alteração do instrumento metodológico consolidado neste documento, "
        "que permanece válido e integralmente aplicável.",
    ]
    for texto in blocos:
        par(doc, texto)


# --------------------------------------------------------------------------
# Capítulo 1 — Introdução e escopo
# --------------------------------------------------------------------------

def capitulo_1(doc):
    h(doc, "1. INTRODUÇÃO E ESCOPO", 1, nova_pagina=True)
    par(
        doc,
        "Este capítulo delimita o objeto do Produto 01, define os termos "
        "técnicos empregados, situa a contratação em seu contexto "
        "institucional, descreve a metodologia e as fontes utilizadas, "
        "declara as limitações que condicionam as conclusões e apresenta a "
        "organização do documento.",
    )

    h(doc, "1.1. Objeto e finalidade do produto", 2)
    par(
        doc,
        "O Produto 01 consiste em documento técnico que diagnostica a "
        "governança dos contratos estratégicos de TIC celebrados pelo "
        "Ministério da Educação com a DATAPREV e com o SERPRO, mediante "
        "análise dos papéis institucionais, das instâncias decisórias, dos "
        "mecanismos de coordenação, dos critérios de monitoramento e dos "
        "riscos associados, acompanhada de matriz analítica e de "
        "recomendações técnicas para o fortalecimento da governança "
        "contratual.",
    )
    par(
        doc,
        "A finalidade do produto é produzir subsídio técnico autônomo e "
        "mensurável para o aperfeiçoamento da gestão contratual de TIC, para "
        "a mitigação de riscos, para a racionalização do gasto público e "
        "para o fortalecimento da tomada de decisão institucional. O produto "
        "não substitui o exercício da gestão e da fiscalização contratual, "
        "que permanecem atribuições dos agentes formalmente designados pelo "
        "órgão, nem constitui manifestação de controle interno ou externo.",
    )
    par(
        doc,
        "O escopo é restrito aos dois contratos identificados. Não integram "
        "o objeto deste produto quaisquer outros contratos, empresas "
        "públicas ou processos administrativos, ainda que relacionados à "
        "área de TIC do órgão.",
    )

    h(doc, "1.2. Delimitação conceitual", 2)
    par(
        doc,
        "Os termos abaixo são empregados neste documento com o sentido "
        "declarado, a fim de evitar ambiguidade na leitura dos achados e "
        "das recomendações.",
    )
    conceitos = [
        ("Governança contratual",
         "Conjunto de estruturas, papéis, instâncias decisórias, mecanismos "
         "de coordenação e controles que asseguram que o contrato seja "
         "dirigido, acompanhado e avaliado de modo a produzir os resultados "
         "pretendidos. Responde à pergunta sobre quem decide, com que "
         "respaldo formal e sob qual controle."),
        ("Gestão contratual",
         "Condução cotidiana da execução do contrato, envolvendo a "
         "autorização de demandas, o acompanhamento dos prazos, a "
         "articulação entre as partes e a adoção das providências "
         "administrativas necessárias à sua execução."),
        ("Fiscalização contratual",
         "Verificação da conformidade entre o que foi pactuado e o que foi "
         "efetivamente executado, com produção de registro apto a sustentar "
         "o recebimento do objeto e o ateste da despesa."),
        ("Instância decisória",
         "Órgão, colegiado ou agente competente para produzir decisão "
         "vinculante sobre matéria relativa ao contrato, no limite de "
         "alçada que lhe tenha sido formalmente atribuído."),
        ("Mecanismo de coordenação",
         "Rotina, canal ou instrumento formal destinado a articular as "
         "atuações do órgão e da contratada, tais como reuniões periódicas, "
         "comitês, fluxos de escalonamento e canais oficiais de "
         "comunicação."),
        ("Achado",
         "Constatação resultante da aplicação de um critério de avaliação "
         "sobre evidência documental identificada. Todo achado deste "
         "produto é classificado como lacuna, sobreposição, fragilidade, "
         "risco ou oportunidade."),
        ("Lacuna",
         "Ausência de elemento de governança exigido por norma, por "
         "instrumento contratual ou pela própria estrutura de gestão "
         "adotada."),
        ("Sobreposição",
         "Atribuição da mesma responsabilidade a mais de um papel, sem "
         "critério de distinção, com risco de duplicidade de esforço ou de "
         "diluição da responsabilidade."),
        ("Fragilidade",
         "Elemento de governança existente, porém insuficiente, "
         "inconsistente ou inefetivo para a finalidade a que se destina."),
        ("Risco",
         "Evento futuro e incerto que, se concretizado, produz efeito "
         "negativo sobre a execução contratual, sobre a conformidade ou "
         "sobre o resultado pretendido."),
        ("Oportunidade",
         "Possibilidade de aprimoramento que não decorre de descumprimento "
         "nem de fragilidade, e cuja adoção eleva a eficiência ou a "
         "maturidade da gestão."),
        ("Não localizado nos autos",
         "Expressão empregada quando o artefato não foi encontrado na "
         "documentação examinada, sem que daí decorra afirmação sobre a sua "
         "inexistência. A afirmação de inexistência somente é feita quando "
         "houver base documental que a sustente."),
        ("Exigência normativa",
         "Obrigação decorrente de norma vigente ou de cláusula contratual "
         "expressa, cujo descumprimento configura desconformidade."),
        ("Boa prática",
         "Recomendação extraída de referencial técnico reconhecido, cuja "
         "adoção é facultativa e cujo descumprimento não configura, por si "
         "só, desconformidade."),
    ]
    tabela(doc, ["Termo", "Sentido adotado neste documento"], conceitos,
           proporcoes=[1, 3], fonte=9)
    nota(
        doc, "Nota.",
        "A distinção entre exigência normativa e boa prática é aplicada de "
        "forma obrigatória em todas as recomendações deste produto, por meio "
        "do campo “Natureza” do quadro do Capítulo 7. A finalidade é impedir "
        "que recomendação de caráter facultativo seja lida como obrigação "
        "legal.",
    )

    h(doc, "1.3. Contexto da contratação", 2)
    par(
        doc,
        "A consultoria foi contratada para produzir três produtos "
        "sequenciais e complementares, todos incidentes sobre os mesmos dois "
        "contratos. O Produto 01 trata da governança contratual; o Produto "
        "02, da análise econômico-financeira; e o Produto 03, da "
        "conformidade normativa e do modelo de monitoramento orientado a "
        "riscos. O quadro a seguir sintetiza a estrutura contratual em que "
        "este produto se insere.",
    )
    tabela(
        doc,
        ["Produto", "Objeto", "Prazo contratual de entrega"],
        [
            ("Produto 01",
             "Diagnóstico da governança contratual, com matriz analítica e "
             "recomendações técnicas", "12/10/2026"),
            ("Produto 02",
             "Análise econômico-financeira, com indicadores, achados "
             "técnicos e recomendações", "20/11/2026"),
            ("Produto 03",
             "Conformidade normativa e modelo de monitoramento orientado a "
             "riscos, com matriz de riscos e critérios de acompanhamento",
             "20/12/2026"),
        ],
        proporcoes=[1, 3.4, 1.2], fonte=9.5,
        alinhamentos=[None, None, "c"],
    )
    par(
        doc,
        "A dependência entre os produtos é de insumo, e não de conclusão. O "
        "mapeamento de papéis e de alçadas realizado neste produto alimenta "
        "a avaliação da segregação de funções no ciclo de pagamento, objeto "
        "do Produto 02, e a identificação de lacunas de controle, objeto do "
        "Produto 03. Cada produto, contudo, revalida à luz do próprio escopo "
        "aquilo que reaproveita, de modo que nenhuma conclusão é transposta "
        "automaticamente de um produto para outro.",
    )
    par(
        doc,
        "O contrato de consultoria indica o Plano Diretor de Tecnologia da "
        "Informação e Comunicação 2025–2027 como referência de alinhamento "
        "institucional. O documento não foi disponibilizado à consultoria "
        "até a data de elaboração deste produto, razão pela qual a "
        "demonstração desse alinhamento permanece pendente e consta da lista "
        "de solicitação documental do Anexo A.",
    )

    h(doc, "1.4. Metodologia e fontes", 2)
    par(
        doc,
        "A metodologia aplicada neste produto está integralmente descrita no "
        "Capítulo 3, que atende à Atividade 1.1. Esta seção registra apenas "
        "os princípios que regem a sua aplicação e as fontes efetivamente "
        "utilizadas na elaboração desta versão.",
    )
    legenda(doc, "Princípios metodológicos aplicados")
    principios = [
        ("Não presunção",
         "Nenhuma informação sobre os contratos é assumida. Toda afirmação "
         "de fato tem origem em documento identificado."),
        ("Rastreabilidade",
         "Todo achado é vinculado ao documento que o fundamenta, com "
         "identificação do processo e da peça no SEI."),
        ("Diferenciação de naturezas",
         "Fato comprovado, interpretação técnica, risco, lacuna de "
         "informação, recomendação e hipótese não validada são registrados "
         "de forma distinta e nunca se confundem."),
        ("Não invenção",
         "Números de processo e de documento, datas, nomes, valores e "
         "identificações normativas jamais são criados, estimados ou "
         "completados por plausibilidade."),
        ("Separação entre exigência e recomendação",
         "Recomendação técnica não é apresentada como obrigação contratual "
         "sem que a exigência esteja comprovada em norma ou em instrumento."),
        ("Independência entre produtos",
         "Cada produto tem escopo próprio; informação anterior é "
         "reaproveitada somente quando pertinente ao escopo em "
         "desenvolvimento."),
        ("Registro de lacunas",
         "O que não puder ser confirmado é registrado como pendência, e não "
         "omitido nem contornado por formulação genérica."),
    ]
    tabela(doc, ["Princípio", "Conteúdo"], principios, proporcoes=[1, 3],
           fonte=9)

    legenda(doc, "Fontes utilizadas na elaboração desta versão")
    fontes = [
        ("F-01", "Contrato de prestação de serviços UNESCO nº %s e seus "
                 "termos de referência" % CONTRATO,
         "Instrumento da consultoria",
         "Fonte primária das exigências, atividades, produtos e prazos"),
        ("F-02", "Documento de orientação do Produto 01",
         "Direcionamento da supervisão",
         "Fonte dos eixos de análise, das atividades e do formato das "
         "matrizes exigidas"),
        ("F-03", "Modelo de referência de produto de outra consultoria",
         "Exemplo de produto",
         "Utilizado exclusivamente quanto à forma: estrutura, formatação e "
         "padrão de apresentação"),
    ]
    tabela(doc, ["ID", "Fonte", "Natureza", "Uso neste produto"], fontes,
           proporcoes=[0.5, 2.4, 1.3, 2.6], fonte=9)
    nota(
        doc, "Ressalva.",
        "A fonte F-03 foi examinada apenas quanto à forma. Nenhum dado, "
        "achado, conclusão, valor, número de contrato ou referência "
        "normativa dela proveniente foi incorporado a este produto. Essa "
        "vedação é absoluta e foi objeto de verificação específica na "
        "revisão final do documento.",
    )
    par(
        doc,
        "Não foi utilizada nenhuma outra fonte. Em especial, não foram "
        "consultados os autos dos Processos SEI nº %s e nº %s, cujas peças "
        "não foram disponibilizadas à consultoria até a data de elaboração "
        "deste documento." % (SEI_DATAPREV, SEI_SERPRO),
    )

    h(doc, "1.5. Limitações", 2)
    par(
        doc,
        "Esta seção declara, de forma expressa e não diluída no restante do "
        "texto, as limitações que condicionam o alcance deste produto.",
    )
    par(
        doc,
        "A limitação principal é de base documental. Até a data de "
        "elaboração deste documento, nenhuma peça dos autos dos dois "
        "processos foi disponibilizada. Em consequência, a caracterização "
        "dos contratos, a identificação dos papéis formalizados, o "
        "mapeamento das instâncias decisórias, a verificação dos mecanismos "
        "de coordenação e a avaliação dos critérios de monitoramento não "
        "puderam ser realizadas. O Capítulo 4 registra essa situação item a "
        "item, e o Anexo A relaciona nominalmente cada documento "
        "solicitado, o motivo da solicitação, a atividade a que se "
        "relaciona e o efeito concreto da sua ausência.",
    )
    par(
        doc,
        "Decorre dessa limitação uma segunda, de natureza normativa. O "
        "documento de orientação do Produto 01 indica referências "
        "normativas cuja identificação oficial está incompleta ou "
        "divergente, conforme detalhado na seção 2.5. Enquanto essa "
        "identificação não for confirmada na fonte oficial, tais "
        "referências não são citadas neste produto com número, ano ou "
        "ementa, por aplicação direta da regra de não invenção.",
    )
    par(
        doc,
        "Há, por fim, uma limitação de delimitação temporal. Não foi "
        "definido o período de execução contratual que serve de recorte à "
        "análise. Essa definição determina o universo documental a examinar "
        "e, enquanto não estabelecida, impede afirmar que a amostra "
        "examinada seja representativa da execução dos contratos.",
    )
    nota(
        doc, "Ressalva.",
        "As limitações aqui declaradas não decorrem de restrição "
        "metodológica nem de opção da consultoria. São condicionantes "
        "externas, integralmente reversíveis mediante a disponibilização "
        "dos documentos relacionados no Anexo A e a resolução das questões "
        "abertas registradas no Anexo C.",
    )

    h(doc, "1.6. Estrutura do documento", 2)
    par(
        doc,
        "O documento organiza-se de modo a espelhar as duas atividades "
        "contratuais que o Produto 01 deve atender, seguido dos "
        "instrumentos que o contrato exige como entregáveis específicos.",
    )
    estrutura = [
        ("Capítulo 1", "Introdução e escopo",
         "Delimita objeto, conceitos, contexto, metodologia e limitações"),
        ("Capítulo 2", "Base normativa e referencial",
         "Relaciona as normas e os referenciais aplicáveis e a situação de "
         "verificação de cada um"),
        ("Capítulo 3", "Atividade 1.1 — Metodologia",
         "Apresenta dimensões, critérios, fontes, normas, referenciais, "
         "escala de severidade e padrões de registro"),
        ("Capítulo 4", "Atividade 1.2 — Diagnóstico",
         "Aplica a metodologia aos quatro eixos de análise e classifica os "
         "achados"),
        ("Capítulo 5", "Matriz analítica",
         "Organiza problema, causa, consequência e melhoria possível, com "
         "evidência e severidade"),
        ("Capítulo 6", "Matriz RACI",
         "Registra a atribuição de responsabilidades apurada e a "
         "recomendada"),
        ("Capítulo 7", "Recomendações técnicas",
         "Vincula cada recomendação a achado identificado, com natureza, "
         "responsável, resultado esperado e indicador"),
        ("Capítulo 8", "Condições para a conclusão do diagnóstico",
         "Consolida dependências documentais, questões abertas e sequência "
         "de retomada"),
        ("Capítulo 9", "Conclusão", "Fecha o argumento do produto"),
        ("Anexos A a D", "Anexos técnicos",
         "Documentos pendentes, instrumentos de coleta, questões e decisões "
         "registradas e referências"),
    ]
    tabela(doc, ["Parte", "Título", "Conteúdo"], estrutura,
           proporcoes=[0.9, 1.8, 3.3], fonte=9)


# --------------------------------------------------------------------------
# Capítulo 2 — Base normativa e referencial
# --------------------------------------------------------------------------

def capitulo_2(doc):
    h(doc, "2. BASE NORMATIVA E REFERENCIAL", 1, nova_pagina=True)
    par(
        doc,
        "Este capítulo relaciona o conjunto normativo e os referenciais "
        "técnicos que sustentam a análise, registrando para cada item a "
        "situação de verificação e a condição de uso. A relação obedece à "
        "regra de não invenção: nenhuma norma é citada com número, ano ou "
        "ementa que não tenha sido confirmado em fonte oficial.",
    )

    h(doc, "2.1. Marco legal", 2)
    par(
        doc,
        "O marco legal aplicável à análise da governança dos contratos "
        "examinados compreende, em primeiro plano, a norma geral de "
        "licitações e contratos administrativos e os princípios "
        "constitucionais que regem a Administração Pública. A Lei nº "
        "14.133, de 2021, interessa a este produto nos blocos relativos às "
        "cláusulas necessárias do contrato, à gestão e à fiscalização "
        "contratual, ao recebimento do objeto e ao regime sancionatório, "
        "por serem esses os pontos em que a norma estabelece exigências "
        "diretamente verificáveis sobre a estrutura de governança. Os "
        "princípios do artigo 37 da Constituição Federal, em especial os da "
        "legalidade, impessoalidade, moralidade, publicidade e eficiência, "
        "operam como parâmetro de avaliação da suficiência dos controles "
        "identificados.",
    )
    par(
        doc,
        "A Lei nº 4.320, de 1964, é registrada neste produto apenas quanto "
        "ao instituto da liquidação da despesa, na medida em que a "
        "existência de lastro documental para o ateste constitui elemento "
        "de governança. O seu desenvolvimento analítico pertence ao escopo "
        "do Produto 02 e não é antecipado aqui.",
    )
    nota(
        doc, "Nota.",
        "A indicação de dispositivos por artigo, inciso ou parágrafo será "
        "incorporada na versão do produto que contiver a análise "
        "documental, quando cada exigência puder ser confrontada com a "
        "cláusula contratual e com a evidência correspondente. Nesta "
        "versão, a referência permanece no nível temático, de modo a não "
        "antecipar vinculação normativa sem o correspondente suporte "
        "probatório.",
    )

    h(doc, "2.2. Marco setorial de TIC", 2)
    par(
        doc,
        "O marco setorial compreende as normas específicas de contratação e "
        "de gestão de soluções de TIC na Administração Pública federal, os "
        "normativos de segurança da informação e as diretrizes "
        "institucionais do órgão. O quadro a seguir registra a situação de "
        "cada item, e o seu conteúdo somente será incorporado à análise "
        "após a confirmação da identificação oficial e a verificação da "
        "aplicabilidade ao objeto efetivamente contratado.",
    )
    legenda(doc, "Marco setorial de TIC e situação de verificação")
    setorial = [
        ("Instrução normativa aplicável à contratação de soluções de TIC "
         "no âmbito do SISP", "Alta", VALID,
         "Confirmar número, ano e vigência em fonte oficial (Q-01)"),
        ("Portarias do SISP com modelos de referência", "A avaliar", VALID,
         "Confirmar identificação e aderência à natureza dos itens "
         "contratados (Q-01)"),
        ("Norma sobre atuação de gestores e fiscais de contrato",
         "Alta", PEND, "Confirmar a norma vigente aplicável ao órgão"),
        ("Normativos de segurança da informação editados pelo GSI",
         "A avaliar", PEND,
         "Mapear os aplicáveis conforme o objeto efetivamente contratado"),
        ("Lei Geral de Proteção de Dados Pessoais", "A avaliar", PEND,
         "Aplicável se houver tratamento de dados pessoais no objeto; "
         "verificar nos autos"),
        ("Plano Diretor de Tecnologia da Informação e Comunicação "
         "2025–2027", "Alta", PEND,
         "Exigido pelo contrato de consultoria como referência de "
         "alinhamento; não disponibilizado (item S-35 do Anexo A)"),
        ("Normativo interno do órgão sobre gestão e fiscalização de "
         "contratos de TIC", "A avaliar", PEND,
         "Verificar existência; item S-36 do Anexo A"),
    ]
    tabela(doc, ["Norma ou instrumento", "Relevância potencial", "Situação",
                 "Providência necessária"],
           setorial, proporcoes=[2.6, 1.1, 1.0, 2.6], fonte=9)

    h(doc, "2.3. Orientações dos órgãos de controle", 2)
    par(
        doc,
        "O documento de orientação do Produto 01 indica expressamente a "
        "consulta a orientações do Tribunal de Contas da União. A Súmula "
        "TCU nº 269 é aplicável a este produto no que diz respeito à "
        "vinculação da remuneração da contratada a resultados e a níveis de "
        "serviço mensuráveis, matéria que integra o Eixo 4 da análise. As "
        "demais indicações constantes do documento de orientação encontram-"
        "se com identificação incompleta e permanecem pendentes de "
        "confirmação.",
    )
    legenda(doc,
            "Orientações de controle e situação de verificação")
    controle = [
        ("Súmula TCU nº 269", "Súmula",
         "Vinculação da remuneração a resultados e a níveis de serviço",
         "Aplicável", "Uso autorizado no Eixo 4"),
        ("“Acórdão 1380/2024”, conforme transcrito no documento de "
         "orientação", "Acórdão",
         "Não identificado o colegiado prolator", VALID,
         "Não citar até a confirmação do órgão julgador, do colegiado e da "
         "ementa em fonte oficial (Q-01)"),
        ("Manifestações de controle relativas especificamente aos "
         "contratos analisados", "Diversa",
         "Apontamentos e diligências eventualmente existentes nos autos",
         PEND,
         "Buscar nos autos dos dois processos (itens S-16 e S-33 do "
         "Anexo A)"),
    ]
    tabela(doc, ["Referência", "Espécie", "Objeto", "Situação",
                 "Condição de uso"],
           controle, proporcoes=[2.0, 0.8, 2.0, 0.9, 2.3], fonte=9)
    nota(
        doc, "Ressalva.",
        "O modelo de referência examinado quanto à forma menciona acórdão "
        "específico relativo a outro par de órgão e empresa. Esse acórdão "
        "não é utilizado neste produto e somente poderá vir a sê-lo se for "
        "localizado de maneira independente nos autos dos processos objeto "
        "desta consultoria, hipótese em que a fonte citada será a dos autos, "
        "e nunca a do modelo.",
    )

    h(doc, "2.4. Referenciais de boas práticas", 2)
    par(
        doc,
        "O documento de orientação determina o uso do COBIT, como "
        "referencial de governança e gestão de TIC, e do ITIL, como "
        "referencial de gestão de serviços de TIC. O COBIT é mobilizado "
        "nesta análise para qualificar achados relativos a "
        "responsabilidades, estrutura decisória, gestão de riscos, "
        "controles e monitoramento. O ITIL é mobilizado para qualificar "
        "achados relativos ao ciclo de demanda, ao tratamento de incidentes "
        "e problemas, à gestão de mudanças, aos níveis de serviço e à "
        "melhoria contínua.",
    )
    par(
        doc,
        "A versão específica de cada referencial a ser adotada e mantida de "
        "forma uniforme nos três produtos ainda não foi definida, razão "
        "pela qual as práticas e os objetivos de governança não são citados "
        "nominalmente nesta versão. A regra que disciplina o emprego dos "
        "referenciais está enunciada na seção 3.6 e é de observância "
        "obrigatória.",
    )

    h(doc, "2.5. Quadro de aplicabilidade", 2)
    par(
        doc,
        "O quadro a seguir consolida a base normativa e referencial do "
        "produto, explicitando, para cada item, a situação de verificação e "
        "a condição que autoriza o seu uso. Itens assinalados como "
        "pendentes de validação não podem ser citados com número, ano ou "
        "ementa enquanto a confirmação não for realizada em fonte oficial.",
    )
    legenda(doc, "Quadro consolidado de aplicabilidade")
    aplic = [
        ("Lei nº 14.133/2021", "Marco legal",
         "Cláusulas necessárias, gestão e fiscalização, recebimento do "
         "objeto e sanções", "Identificação confirmada",
         "Aplicável; citação por dispositivo na etapa de análise "
         "documental"),
        ("Constituição Federal, art. 37", "Marco legal",
         "Princípios da Administração Pública", "Identificação confirmada",
         "Aplicável como parâmetro de suficiência dos controles"),
        ("Lei nº 4.320/1964", "Marco legal",
         "Liquidação da despesa", "Identificação confirmada",
         "Aplicável de forma restrita; desenvolvimento no Produto 02"),
        ("Norma do SISP sobre contratação de soluções de TIC",
         "Marco setorial", "Processo de contratação e gestão de TIC", VALID,
         "Vedada a citação até a confirmação oficial (Q-01)"),
        ("“Portaria 5.950”, conforme transcrito no documento de orientação",
         "Marco setorial", "Não informado o órgão emissor nem o ano", VALID,
         "Vedada a citação até a confirmação oficial (Q-01)"),
        ("“IN 94/2023”, conforme transcrito no documento de orientação",
         "Marco setorial",
         "Ano possivelmente divergente da norma efetivamente aplicável",
         VALID, "Vedada a citação até a confirmação oficial (Q-01)"),
        ("Súmula TCU nº 269", "Orientação de controle",
         "Vinculação da remuneração a resultados e níveis de serviço",
         "Identificação confirmada", "Aplicável no Eixo 4"),
        ("“Acórdão 1380/2024”, conforme transcrito no documento de "
         "orientação", "Orientação de controle",
         "Colegiado prolator não informado", VALID,
         "Vedada a citação até a confirmação oficial (Q-01)"),
        ("PDTIC 2025–2027", "Diretriz institucional",
         "Alinhamento institucional exigido pelo contrato de consultoria",
         PEND, "Documento não disponibilizado (item S-35 do Anexo A)"),
        ("COBIT", "Referencial de boas práticas",
         "Governança e gestão de TIC, responsabilidades, riscos, controles "
         "e monitoramento", "Versão a definir (Q-02)",
         "Uso restrito à qualificação de achados e à fundamentação de "
         "recomendações"),
        ("ITIL", "Referencial de boas práticas",
         "Gestão de serviços, incidentes, problemas, mudanças e níveis de "
         "serviço", "Versão a definir (Q-02)",
         "Uso restrito à qualificação de achados e à fundamentação de "
         "recomendações"),
    ]
    tabela(doc, ["Referência", "Categoria", "Objeto de interesse",
                 "Situação", "Condição de uso"],
           aplic, proporcoes=[1.9, 1.1, 2.1, 1.2, 2.1], fonte=8.5)
    nota(
        doc, "Nota.",
        "As referências transcritas entre aspas reproduzem literalmente a "
        "forma como constam do documento de orientação do Produto 01. A "
        "transcrição não equivale a citação normativa e não afirma a "
        "existência, a vigência ou o conteúdo das normas assim designadas. "
        "A confirmação da identificação oficial constitui a questão aberta "
        "Q-01, registrada no Anexo C.",
    )


def fluxo(doc, etapas, correto=True):
    """Representa uma sequência metodológica em linha destacada."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(10)
    cor = AZUL if correto else RGBColor(0x8A, 0x3A, 0x1F)
    for i, etapa in enumerate(etapas):
        if i:
            sep = p.add_run("   →   ")
            sep.font.size = Pt(10)
            sep.font.color.rgb = CINZA
            sep.font.name = FONTE
        r = p.add_run(etapa)
        r.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = cor
        r.font.name = FONTE
    b = _el("w:pBdr")
    for lado in ("top", "bottom"):
        b.append(_el("w:" + lado, val="single", sz="4", space="6",
                     color="B7C0D0"))
    p._p.get_or_add_pPr().append(b)
    return p


# --------------------------------------------------------------------------
# Capítulo 3 — Atividade 1.1
# --------------------------------------------------------------------------

def capitulo_3(doc):
    h(doc,
      "3. ATIVIDADE 1.1 — METODOLOGIA DE ANÁLISE DA GOVERNANÇA CONTRATUAL",
      1, nova_pagina=True)
    par(
        doc,
        "Este capítulo atende integralmente à Atividade 1.1 do contrato de "
        "consultoria, que exige a definição da metodologia de análise da "
        "governança contratual, contemplando critérios de avaliação, "
        "dimensões de análise, fontes documentais e parâmetros normativos "
        "aplicáveis aos contratos estratégicos de TIC. A metodologia aqui "
        "definida é o instrumento único de produção dos achados do Capítulo "
        "4 e permanece aplicável independentemente do momento em que a base "
        "documental venha a ser disponibilizada.",
    )

    h(doc, "3.1. Sequência metodológica adotada", 2)
    par(
        doc,
        "O documento de orientação do Produto 01 estabelece a sequência que "
        "ordena a construção do instrumento analítico. Essa sequência é "
        "adotada na íntegra e organiza as seções seguintes deste capítulo.",
    )
    fluxo(doc, ["Critérios", "Dimensões", "Documentos", "Normas",
                "Referenciais"])
    par(
        doc,
        "A ordem não é meramente expositiva. Ela impede que a análise "
        "comece pelo referencial técnico e termine na busca de evidência "
        "que o confirme, inversão que produziria diagnóstico aparentemente "
        "fundamentado, porém descolado da realidade dos contratos "
        "examinados. A sequência assegura que o critério anteceda a "
        "evidência, que a evidência anteceda o achado e que o achado "
        "anteceda a recomendação.",
    )
    legenda(doc, "Elementos da metodologia e seção correspondente")
    tabela(
        doc,
        ["Elemento exigido", "O que define", "Seção"],
        [
            ("Critérios", "O que será avaliado em cada contrato", "3.3"),
            ("Dimensões", "Sob que perspectiva a avaliação é feita", "3.2"),
            ("Documentos", "Onde a evidência é buscada", "3.4"),
            ("Normas", "Qual exigência serve de parâmetro", "3.5"),
            ("Referenciais",
             "Que prática reconhecida explica por que o achado é um "
             "problema", "3.6"),
        ],
        proporcoes=[1.2, 3.6, 0.7], fonte=9.5,
        alinhamentos=[None, None, "c"],
    )

    h(doc, "3.2. Dimensões de análise", 2)
    par(
        doc,
        "A análise opera em duas dimensões, ambas exigidas pelo documento "
        "de orientação. A dimensão estratégica examina as decisões, as "
        "estruturas e os alinhamentos de nível institucional. A dimensão "
        "operacional examina a execução cotidiana da gestão contratual. "
        "Cada dimensão é decomposta em eixos internos, que operacionalizam "
        "a sua aplicação e vinculam-na aos critérios do item 3.3.",
    )

    h(doc, "3.2.1. Dimensão estratégica", 3)
    par(
        doc,
        "A dimensão estratégica responde à pergunta sobre como o contrato é "
        "dirigido. Interessa-lhe a existência de estrutura formal de "
        "decisão, a cobertura dos papéis exigidos, a clareza das alçadas e "
        "o tratamento institucional dos riscos.",
    )
    legenda(doc, "Eixos internos da dimensão estratégica")
    tabela(
        doc,
        ["Eixo interno", "Objeto de exame", "Critérios vinculados"],
        [
            ("Alinhamento institucional",
             "Vinculação dos contratos a diretrizes e planos institucionais",
             "E-11"),
            ("Estrutura decisória",
             "Quem decide o quê, em que nível e com qual respaldo formal",
             "E-09, E-10"),
            ("Papéis institucionais",
             "Designação formal, descrição de atribuições e cobertura dos "
             "papéis exigidos", "E-01 a E-07"),
            ("Governança colegiada",
             "Existência, composição e funcionamento de instância "
             "colegiada", "E-08"),
            ("Gestão estratégica de riscos",
             "Identificação, tratamento e escalonamento de riscos em nível "
             "institucional", "E-12"),
        ],
        proporcoes=[1.4, 3.3, 1.1], fonte=9,
    )

    h(doc, "3.2.2. Dimensão operacional", 3)
    par(
        doc,
        "A dimensão operacional responde à pergunta sobre como o contrato é "
        "executado e acompanhado no dia a dia. Interessa-lhe o ciclo "
        "completo da demanda, desde a solicitação até o arquivamento da "
        "evidência que sustenta o ateste.",
    )
    legenda(doc, "Eixos internos da dimensão operacional")
    tabela(
        doc,
        ["Eixo interno", "Objeto de exame", "Critérios vinculados"],
        [
            ("Ciclo da demanda",
             "Como uma Ordem de Serviço nasce, é aprovada, executada e "
             "encerrada", "O-01, O-02, O-03"),
            ("Comunicação e coordenação",
             "Canais, periodicidade e formalização da interação entre órgão "
             "e contratada", "O-04, O-05, O-06"),
            ("Fiscalização e acompanhamento",
             "Como a execução é acompanhada, atestada e recebida",
             "O-09, O-10, O-11"),
            ("Monitoramento e indicadores",
             "Quais indicadores existem e como são apurados",
             "O-07, O-08"),
            ("Tratamento de incidentes e problemas",
             "Registro, escalonamento e resolução de ocorrências", "O-12"),
            ("Registro documental",
             "Onde e como decisões e evidências são arquivadas",
             "O-13, O-14"),
        ],
        proporcoes=[1.4, 3.3, 1.1], fonte=9,
    )
    nota(
        doc, "Nota.",
        "A correspondência entre eixos internos e critérios é vinculante. "
        "Nenhum achado é produzido fora de um critério, e nenhum critério "
        "existe sem eixo que o justifique. Essa amarração é o que permite "
        "auditar a completude do diagnóstico ao final da aplicação.",
    )

    h(doc, "3.3. Critérios de avaliação", 2)
    par(
        doc,
        "Os critérios constituem o instrumento de aplicação das dimensões. "
        "Cada critério recebe, na etapa de análise documental, uma situação "
        "apurada segundo a escala definida e a indicação da fonte que a "
        "sustenta. São vinte e seis critérios, doze na dimensão estratégica "
        "e catorze na operacional, aplicados separadamente a cada um dos "
        "dois contratos.",
    )

    h(doc, "3.3.1. Critérios da dimensão estratégica", 3)
    legenda(doc, "Critérios da dimensão estratégica")
    tabela(
        doc,
        ["ID", "Critério", "Escala de avaliação"],
        [
            ("E-01", "Gestor do contrato formalmente designado",
             "Formalizado · Parcial · Não localizado"),
            ("E-02", "Fiscal técnico formalmente designado",
             "Formalizado · Parcial · Não localizado"),
            ("E-03", "Fiscal administrativo formalmente designado",
             "Formalizado · Parcial · Não localizado"),
            ("E-04", "Fiscal requisitante ou setorial designado",
             "Formalizado · Parcial · Não localizado · Não aplicável"),
            ("E-05", "Substitutos designados para os papéis críticos",
             "Formalizado · Parcial · Não localizado"),
            ("E-06", "Gestores das soluções identificados por item "
                     "contratado",
             "Formalizado · Parcial · Não localizado"),
            ("E-07", "Atribuições descritas nos atos de designação",
             "Detalhadas · Genéricas · Ausentes"),
            ("E-08", "Instância colegiada de governança instituída",
             "Instituída · Informal · Inexistente"),
            ("E-09", "Níveis de alçada decisória definidos",
             "Definidos · Parciais · Indefinidos"),
            ("E-10",
             "Participação de agentes não formalizados nas decisões",
             "Não identificada · Identificada · Indeterminado"),
            ("E-11",
             "Alinhamento declarado a plano ou diretriz institucional",
             "Explícito · Implícito · Ausente"),
            ("E-12", "Mapa ou matriz de riscos do contrato existente",
             "Existente e atualizado · Existente e desatualizado · "
             "Inexistente"),
        ],
        proporcoes=[0.55, 2.9, 2.4], fonte=9,
        alinhamentos=["c", None, None],
    )

    h(doc, "3.3.2. Critérios da dimensão operacional", 3)
    legenda(doc, "Critérios da dimensão operacional")
    tabela(
        doc,
        ["ID", "Critério", "Escala de avaliação"],
        [
            ("O-01", "Fluxo de solicitação de Ordem de Serviço documentado",
             "Documentado · Praticado sem documentação · Indefinido"),
            ("O-02",
             "Fluxo de aprovação de Ordem de Serviço com alçada definida",
             "Definido · Parcial · Indefinido"),
            ("O-03", "Registro centralizado das demandas",
             "Centralizado · Disperso · Inexistente"),
            ("O-04",
             "Reuniões de acompanhamento com periodicidade definida",
             "Regular · Irregular · Inexistente"),
            ("O-05", "Atas ou registros formais das reuniões",
             "Sistemáticos · Esporádicos · Inexistentes"),
            ("O-06", "Canal formal de comunicação com a contratada",
             "Formalizado · Informal · Indefinido"),
            ("O-07", "Indicadores de desempenho definidos",
             "Pactuados · Parcialmente pactuados · Não pactuados"),
            ("O-08", "Fonte de apuração dos indicadores",
             "Independente · Dependente da contratada · Inexistente"),
            ("O-09", "Relatórios de fiscalização emitidos",
             "Sistemáticos · Pontuais · Não localizados"),
            ("O-10", "Termos de recebimento emitidos",
             "Sistemáticos · Pontuais · Não localizados"),
            ("O-11", "Ateste vinculado a evidência de verificação",
             "Vinculado · Formal sem evidência · Não localizado"),
            ("O-12", "Procedimento de escalonamento de incidentes",
             "Formalizado · Informal · Inexistente"),
            ("O-13",
             "Trilha de evidências arquivada e vinculada à demanda",
             "Íntegra · Parcial · Ausente"),
            ("O-14",
             "Normas de cumprimento obrigatório mapeadas no contrato",
             "Mapeadas · Parciais · Não mapeadas"),
        ],
        proporcoes=[0.55, 2.9, 2.4], fonte=9,
        alinhamentos=["c", None, None],
    )

    h(doc, "3.3.3. Regras de aplicação das escalas", 3)
    par(
        doc,
        "A atribuição de situação apurada obedece a três regras de "
        "observância obrigatória, destinadas a impedir que a limitação de "
        "acesso documental se converta em afirmação sobre a realidade dos "
        "contratos.",
    )
    par(
        doc,
        "A primeira regra diz respeito ao emprego da expressão “não "
        "localizado”. Ela é utilizada sempre que a ausência do artefato "
        "puder decorrer de limitação de acesso à documentação, e não de "
        "ausência real. A expressão “inexistente” somente é empregada "
        "quando houver base documental que sustente a afirmação, como "
        "manifestação expressa do órgão nesse sentido ou constatação "
        "registrada em peça dos autos.",
    )
    par(
        doc,
        "A segunda regra impõe a indicação da fonte. Nenhum critério recebe "
        "situação apurada sem a identificação do documento que a sustenta, "
        "com processo, peça e localização interna quando aplicável. "
        "Critério sem fonte permanece sem avaliação e é registrado como "
        "lacuna de informação, na forma do item 3.8.",
    )
    par(
        doc,
        "A terceira regra determina a avaliação separada por contrato. Cada "
        "critério é avaliado individualmente para a DATAPREV e para o "
        "SERPRO, ainda que o resultado seja idêntico. A consolidação "
        "comparativa é produto da análise, e não premissa dela.",
    )

    h(doc, "3.4. Fontes documentais", 2)
    par(
        doc,
        "A evidência que sustenta os achados é buscada nas nove categorias "
        "documentais abaixo. A relação é exaustiva quanto às categorias e "
        "indicativa quanto às peças, já que o conjunto efetivamente "
        "existente somente pode ser conhecido após o exame dos autos.",
    )
    legenda(doc, "Categorias de fontes documentais")
    tabela(
        doc,
        ["Categoria", "Peças compreendidas", "Eixos que alimenta"],
        [
            ("Instrumentos contratuais",
             "Contratos, termos aditivos e apostilamentos",
             "Todos"),
            ("Instrumentos de planejamento",
             "Estudo Técnico Preliminar, Termo de Referência e demais "
             "artefatos de planejamento da contratação", "Eixos 1, 2 e 4"),
            ("Atos de designação",
             "Portarias, ordens de serviço internas e atos de nomeação de "
             "gestores e fiscais", "Eixo 1"),
            ("Execução contratual",
             "Ordens de Serviço, propostas de atendimento, notas fiscais e "
             "atestes", "Eixos 3 e 4"),
            ("Fiscalização",
             "Relatórios de fiscalização, termos de recebimento provisório "
             "e definitivo e listas de verificação", "Eixo 4"),
            ("Governança",
             "Atas de reunião, regimentos de comitês, matrizes de "
             "responsabilidade e fluxos formalizados", "Eixos 2 e 3"),
            ("Riscos",
             "Mapas de gerenciamento de riscos, registros de incidentes e "
             "escalonamentos", "Eixo 4"),
            ("Controle",
             "Manifestações de órgãos de controle, respostas a diligências "
             "e planos de providências", "Todos"),
            ("Institucional",
             "Plano diretor de TIC, normativos internos e organogramas",
             "Eixos 1 e 2"),
        ],
        proporcoes=[1.4, 3.4, 1.1], fonte=9,
    )
    par(
        doc,
        "A relação nominal das peças solicitadas, com motivo, atividade "
        "relacionada, efeito da ausência e alternativa de comprovação, "
        "consta do Anexo A. Nenhuma categoria acima foi acessada até a data "
        "de elaboração deste documento.",
    )

    h(doc, "3.5. Parâmetros normativos aplicáveis", 2)
    par(
        doc,
        "Os parâmetros normativos que servem de referência à avaliação "
        "estão relacionados no Capítulo 2, com a respectiva situação de "
        "verificação. Esta seção registra a regra que disciplina o seu "
        "emprego na produção dos achados.",
    )
    par(
        doc,
        "Norma somente é invocada como parâmetro quando três condições "
        "estiverem simultaneamente satisfeitas: a identificação oficial "
        "estiver confirmada em fonte oficial, quanto a número, ano e "
        "ementa; a aplicabilidade ao objeto efetivamente contratado estiver "
        "verificada nos autos; e a exigência dela decorrente for passível "
        "de confronto com evidência documental determinada. Ausente "
        "qualquer das condições, a norma não é citada, e a questão é "
        "registrada como pendência.",
    )
    nota(
        doc, "Nota.",
        "Esta regra explica por que o Capítulo 2 relaciona referências "
        "assinaladas como pendentes de validação sem delas extrair "
        "exigência. Relacionar a referência é necessário para documentar o "
        "que foi indicado pela supervisão e o que precisa ser confirmado; "
        "extrair exigência antes da confirmação seria fundamentar achado em "
        "norma cuja identidade não se conhece.",
    )

    h(doc, "3.6. Referenciais de boas práticas e regra de uso", 2)
    par(
        doc,
        "Os referenciais de boas práticas cumprem função delimitada nesta "
        "metodologia: fundamentam recomendações e qualificam achados, mas "
        "não os geram. A sequência de raciocínio admitida é uma só.",
    )
    fluxo(doc, ["Evidência documental", "Achado",
                "Referencial que explica o problema", "Recomendação"])
    par(doc, "A sequência inversa é expressamente vedada.")
    fluxo(doc, ["Referencial", "Recomendação",
                "Busca de evidência que a justifique"], correto=False)
    par(
        doc,
        "A inversão produz documento que parece fundamentado e não é. Ela "
        "transforma o referencial em premissa, e não em instrumento de "
        "leitura, e conduz à recomendação de práticas que podem não "
        "corresponder a nenhum problema real dos contratos examinados. A "
        "vedação é absoluta e integra a lista de verificação aplicada ao "
        "fechamento do produto.",
    )
    par(
        doc,
        "A correspondência entre os eixos de análise e as práticas "
        "específicas de cada referencial será estabelecida quando a versão "
        "a adotar estiver definida, conforme a questão aberta Q-02. O "
        "quadro a seguir registra a estrutura dessa correspondência e o seu "
        "estado atual.",
    )
    legenda(doc,
            "Correspondência prevista entre eixos e referenciais")
    tabela(
        doc,
        ["Eixo de análise", "COBIT — domínio ou objetivo",
         "ITIL — prática", "Situação"],
        [
            ("Papéis e responsabilidades", PEND, PEND, "A definir (Q-02)"),
            ("Instâncias decisórias", PEND, PEND, "A definir (Q-02)"),
            ("Comunicação e coordenação", PEND, PEND, "A definir (Q-02)"),
            ("Monitoramento e níveis de serviço", PEND, PEND,
             "A definir (Q-02)"),
            ("Gestão de riscos", PEND, PEND, "A definir (Q-02)"),
            ("Incidentes e problemas", PEND, PEND, "A definir (Q-02)"),
        ],
        proporcoes=[2.0, 1.7, 1.3, 1.3], fonte=9,
        alinhamentos=[None, "c", "c", "c"],
    )

    h(doc, "3.7. Escala de severidade dos achados", 2)
    par(
        doc,
        "Todo achado recebe severidade, atribuída após a sua identificação "
        "e sempre acompanhada da justificativa do enquadramento. A "
        "severidade qualifica o achado; não o cria.",
    )
    legenda(doc, "Escala de severidade")
    tabela(
        doc,
        ["Severidade", "Definição", "Critério de enquadramento"],
        [
            ("Alta",
             "Descumprimento de exigência normativa ou contratual expressa, "
             "ou ausência de controle essencial",
             "Risco de responsabilização, de desconformidade ou de prejuízo "
             "relevante"),
            ("Média",
             "Fragilidade que compromete a efetividade da gestão sem "
             "configurar descumprimento direto",
             "Risco operacional relevante, porém contornável"),
            ("Baixa",
             "Oportunidade de aprimoramento sem risco imediato",
             "Ganho de eficiência ou de maturidade da gestão"),
        ],
        proporcoes=[0.9, 2.8, 2.4], fonte=9,
        alinhamentos=["c", None, None],
    )

    h(doc, "3.8. Padrões de registro e rastreabilidade", 2)
    par(
        doc,
        "Cada natureza de conteúdo é registrada em formato próprio. A "
        "padronização impede que interpretação seja lida como fato e que "
        "hipótese se converta em conclusão ao longo das etapas do trabalho. "
        "Os formatos abaixo são aplicados na etapa de análise documental e "
        "estão reproduzidos, em versão de preenchimento, no Anexo B.",
    )
    legenda(doc, "Padrões de registro por natureza de conteúdo")
    tabela(
        doc,
        ["Natureza", "Campos obrigatórios", "Finalidade do registro"],
        [
            ("Fato comprovado",
             "Descrição; fonte, com processo, peça e data; localização "
             "interna no documento",
             "Fixar o que o documento estabelece, sem juízo"),
            ("Interpretação técnica",
             "Afirmação; fatos que a sustentam, com fontes; fundamento "
             "técnico ou normativo; grau de certeza",
             "Explicitar o raciocínio e o seu limite de confiabilidade"),
            ("Risco identificado",
             "Descrição; causa raiz; evidência; probabilidade e impacto, "
             "ambos justificados; controle existente e sua efetividade",
             "Sustentar a avaliação de risco em evidência, e não em "
             "conjectura"),
            ("Lacuna de informação",
             "O que falta; por que é necessária; o que fica impedido sem "
             "ela; alternativa de comprovação; situação",
             "Tornar a ausência visível e administrável"),
            ("Recomendação técnica",
             "Problema que endereça; ação proposta; natureza; fundamento; "
             "responsável sugerido, por papel; resultado esperado; "
             "indicador de implementação; prazo sugerido",
             "Garantir que a recomendação seja verificável e vinculada a "
             "achado"),
        ],
        proporcoes=[1.1, 2.7, 2.3], fonte=9,
    )
    nota(
        doc, "Nota.",
        "No registro de recomendações, o campo “Natureza” e o campo "
        "“Responsável” são de preenchimento obrigatório. O primeiro impede "
        "que boa prática seja apresentada como obrigação legal. O segundo "
        "é preenchido com o papel, e nunca com o nome da pessoa, salvo "
        "quando o produto exigir identificação nominal e houver base "
        "documental que a sustente.",
    )
    par(
        doc,
        "A rastreabilidade é assegurada por identificadores estáveis. "
        "Achados recebem prefixo conforme a categoria — L para lacunas, S "
        "para sobreposições, F para fragilidades, R para riscos e OP para "
        "oportunidades. Itens da matriz analítica recebem o prefixo MA e "
        "recomendações, o prefixo REC. Toda recomendação cita, "
        "obrigatoriamente, ao menos um identificador de achado, e nenhum "
        "achado é enunciado sem a fonte documental que o sustenta.",
    )
    legenda(doc, "Sistema de identificadores")
    tabela(
        doc,
        ["Prefixo", "Aplica-se a", "Vinculação obrigatória"],
        [
            ("L-nn", "Lacunas", "Evidência documental"),
            ("S-nn", "Sobreposições", "Evidência documental"),
            ("F-nn", "Fragilidades", "Evidência documental"),
            ("R-nn", "Riscos", "Evidência documental e causa raiz"),
            ("OP-nn", "Oportunidades", "Achado ou evidência de contexto"),
            ("MA-nn", "Itens da matriz analítica",
             "Ao menos um achado de origem"),
            ("REC-nn", "Recomendações",
             "Ao menos um item da matriz analítica"),
        ],
        proporcoes=[0.9, 2.2, 2.6], fonte=9,
        alinhamentos=["c", None, None],
    )

    h(doc, "3.9. Procedimento de aplicação", 2)
    par(
        doc,
        "A metodologia é aplicada na sequência abaixo. A ordem é "
        "vinculante: nenhuma etapa se inicia antes da conclusão da "
        "anterior, e o produto de cada etapa é insumo verificável da "
        "seguinte.",
    )
    legenda(doc, "Etapas de aplicação da metodologia")
    tabela(
        doc,
        ["Etapa", "Atividade", "Produto da etapa"],
        [
            ("1", "Recebimento e catalogação das peças dos autos, com "
                  "preenchimento da ficha de documento analisado",
             "Registro de documentos analisados (Anexo B, seção B.4)"),
            ("2", "Caracterização dos contratos e do modelo de gestão "
                  "pactuado", "Seções 4.2.1 a 4.2.3 preenchidas"),
            ("3", "Aplicação dos critérios por contrato, com indicação de "
                  "fonte", "Situação apurada para os vinte e seis "
                  "critérios, em cada contrato"),
            ("4", "Confronto entre o previsto e o praticado, por eixo de "
                  "análise", "Seções 4.3 a 4.6 preenchidas"),
            ("5", "Classificação dos achados em lacunas, sobreposições, "
                  "fragilidades, riscos e oportunidades",
             "Seção 4.7 preenchida, com severidade justificada"),
            ("6", "Consolidação da matriz analítica",
             "Capítulo 5 preenchido"),
            ("7", "Consolidação das matrizes de responsabilidade, apurada "
                  "e recomendada", "Capítulo 6 preenchido"),
            ("8", "Formulação das recomendações, vinculadas por "
                  "identificador", "Capítulo 7 preenchido"),
            ("9", "Revisão técnica e aplicação da lista de verificação de "
                  "fechamento", "Produto validado para emissão"),
        ],
        proporcoes=[0.6, 3.2, 2.4], fonte=9,
        alinhamentos=["c", None, None],
    )

    h(doc, "3.10. Quadro-síntese da metodologia", 2)
    par(
        doc,
        "O quadro a seguir consolida o instrumento analítico definido neste "
        "capítulo e encerra a Atividade 1.1.",
    )
    legenda(doc, "Síntese do instrumento analítico")
    tabela(
        doc,
        ["Componente", "Definição adotada", "Situação"],
        [
            ("Dimensões de análise",
             "Duas dimensões, estratégica e operacional, decompostas em "
             "onze eixos internos", "Definida"),
            ("Critérios de avaliação",
             "Vinte e seis critérios com escala própria e regra de "
             "atribuição", "Definida"),
            ("Fontes documentais",
             "Nove categorias, com relação nominal de peças no Anexo A",
             "Definida"),
            ("Parâmetros normativos",
             "Marco legal confirmado; marco setorial e orientações de "
             "controle parcialmente pendentes de validação oficial",
             "Parcial (Q-01)"),
            ("Referenciais de boas práticas",
             "COBIT e ITIL, com regra de uso definida e versão a fixar",
             "Parcial (Q-02)"),
            ("Escala de severidade",
             "Três níveis, com definição e critério de enquadramento",
             "Definida"),
            ("Padrões de registro",
             "Cinco formatos, por natureza de conteúdo, com campos "
             "obrigatórios", "Definida"),
            ("Sistema de identificadores",
             "Sete prefixos, com vinculação obrigatória entre achado, "
             "matriz e recomendação", "Definida"),
            ("Procedimento de aplicação",
             "Nove etapas sequenciais, com produto verificável por etapa",
             "Definida"),
        ],
        proporcoes=[1.4, 3.4, 1.1], fonte=9,
    )
    nota(
        doc, "Nota.",
        "Os dois componentes assinalados como parciais dependem "
        "exclusivamente da resolução das questões abertas Q-01 e Q-02, "
        "registradas no Anexo C. Ambas são externas à modelagem "
        "metodológica e não afetam a validade dos demais componentes, que "
        "estão aptos a aplicação imediata.",
    )


# --------------------------------------------------------------------------
# Capítulo 4 — Atividade 1.2
# --------------------------------------------------------------------------

def _bloco_contrato(doc, nome, processo, ordem):
    h(doc, "4.2.%d. Contrato %s" % (ordem, nome), 3)
    par(
        doc,
        "Processo SEI nº %s. Os campos abaixo serão preenchidos "
        "exclusivamente a partir do instrumento contratual e dos artefatos "
        "de planejamento correspondentes, quando disponibilizados."
        % processo,
    )
    legenda(doc, "Identificação do contrato %s" % nome)
    tabela(
        doc,
        ["Campo", "Informação", "Situação"],
        [(c, "—", PEND) for c in [
            "Número do contrato", "Objeto contratual",
            "Fundamento legal da contratação", "Data de assinatura",
            "Vigência", "Valor global", "Aditivos celebrados",
            "Termo de Referência aplicável",
            "Estudo Técnico Preliminar", "Itens contratados",
            "Métrica de faturamento por item",
        ]],
        proporcoes=[2.0, 2.6, 1.0], fonte=9,
        alinhamentos=[None, "c", "c"],
    )
    legenda(doc, "Papéis institucionais formalizados no contrato %s"
            % nome)
    tabela(
        doc,
        ["Papel", "Nome", "Ato de designação", "Data", "Situação"],
        [(p, "—", "—", "—", PEND) for p in [
            "Gestor do contrato", "Gestor substituto", "Fiscal técnico",
            "Fiscal técnico substituto", "Fiscal administrativo",
            "Fiscal requisitante ou setorial",
            "Gestores das soluções, por item", "Preposto da contratada",
        ]],
        proporcoes=[2.2, 1.2, 1.6, 0.8, 1.0], fonte=9,
        alinhamentos=[None, "c", "c", "c", "c"],
    )
    legenda(doc, "Instâncias e mecanismos de governança do contrato %s"
            % nome)
    tabela(
        doc,
        ["Elemento", "Situação apurada", "Fonte", "Situação"],
        [(e, "—", "—", PEND) for e in [
            "Comitê ou instância colegiada de acompanhamento",
            "Periodicidade das reuniões",
            "Forma de registro das decisões",
            "Canal formal de comunicação entre órgão e contratada",
            "Fluxo de solicitação e aprovação de Ordens de Serviço",
            "Fluxo de escalonamento de problemas e incidentes",
            "Matriz de responsabilidades formalizada",
        ]],
        proporcoes=[2.6, 1.6, 1.0, 1.0], fonte=9,
        alinhamentos=[None, "c", "c", "c"],
    )


def _quadro_eixo(doc, numero, titulo, descricao, pontos):
    h(doc, "4.%d. Eixo %d — %s" % (numero + 2, numero, titulo), 2)
    par(doc, descricao)
    legenda(doc, "Eixo %d: pontos de verificação" % numero)
    tabela(
        doc,
        ["Ponto de verificação", "DATAPREV", "SERPRO",
         "Fonte da evidência", "Situação"],
        [(p, "—", "—", "—", PEND) for p in pontos],
        proporcoes=[2.6, 0.9, 0.9, 1.4, 1.0], fonte=9,
        alinhamentos=[None, "c", "c", "c", "c"],
    )


def capitulo_4(doc):
    h(doc, "4. ATIVIDADE 1.2 — DIAGNÓSTICO DA GOVERNANÇA CONTRATUAL", 1,
      nova_pagina=True)
    par(
        doc,
        "Este capítulo aplica aos dois contratos o instrumento analítico "
        "definido no Capítulo 3 e organiza o registro dos achados nas cinco "
        "categorias exigidas pelo documento de orientação: lacunas, "
        "sobreposições, fragilidades, riscos e oportunidades. A estrutura "
        "de registro está integralmente montada; o preenchimento depende da "
        "base documental cuja situação a seção 4.1 declara.",
    )

    h(doc, "4.1. Situação da base documental", 2)
    par(
        doc,
        "Até a data de elaboração deste documento, nenhuma peça dos autos "
        "dos Processos SEI nº %s e nº %s foi disponibilizada à consultoria. "
        "O quadro a seguir registra a situação por categoria documental e "
        "o efeito concreto da ausência sobre a análise."
        % (SEI_DATAPREV, SEI_SERPRO),
    )
    legenda(doc, "Situação da base documental por categoria")
    tabela(
        doc,
        ["Categoria documental", "Situação", "Análise impedida"],
        [
            ("Instrumentos contratuais", "Não disponibilizada",
             "Caracterização do objeto, do modelo de gestão e das "
             "obrigações de fiscalização pactuadas"),
            ("Instrumentos de planejamento", "Não disponibilizada",
             "Confronto entre o modelo de gestão concebido e o praticado"),
            ("Atos de designação", "Não disponibilizada",
             "Eixo 1, na íntegra, e matriz de responsabilidades apurada"),
            ("Execução contratual", "Não disponibilizada",
             "Ciclo da demanda e verificação do fluxo real de aprovação"),
            ("Fiscalização", "Não disponibilizada",
             "Efetividade do acompanhamento e rito de recebimento"),
            ("Governança", "Não disponibilizada",
             "Eixos 2 e 3, quanto a instâncias, coordenação e registro de "
             "decisões"),
            ("Riscos", "Não disponibilizada",
             "Identificação e tratamento formal de riscos"),
            ("Controle", "Não disponibilizada",
             "Verificação de apontamentos já existentes sobre os "
             "contratos"),
            ("Institucional", "Não disponibilizada",
             "Alinhamento institucional e mapeamento dos níveis de "
             "decisão"),
        ],
        proporcoes=[1.6, 1.3, 3.1], fonte=9,
        alinhamentos=[None, "c", None],
    )
    nota(
        doc, "Ressalva.",
        "A expressão “não disponibilizada” descreve a situação de acesso da "
        "consultoria, e não a existência das peças nos autos. Nenhuma "
        "afirmação deste capítulo deve ser lida como constatação de "
        "ausência de artefato no âmbito do órgão. A distinção entre não "
        "localizado e inexistente, definida na seção 3.3.3, é aplicada "
        "integralmente.",
    )

    h(doc, "4.2. Caracterização dos contratos analisados", 2)
    par(
        doc,
        "A caracterização precede a avaliação. Sem conhecer o objeto "
        "contratado, os itens que o compõem e a métrica de faturamento de "
        "cada um, não é possível determinar quais elementos de governança "
        "são exigíveis nem qual estrutura de fiscalização seria adequada. "
        "Os quadros desta seção constituem, por isso, pré-requisito dos "
        "eixos 1 a 4.",
    )
    _bloco_contrato(doc, "DATAPREV", SEI_DATAPREV, 1)
    _bloco_contrato(doc, "SERPRO", SEI_SERPRO, 2)

    h(doc, "4.2.3. Quadro comparativo entre os contratos", 3)
    par(
        doc,
        "O quadro comparativo sustenta a análise transversal exigida nos "
        "três produtos da consultoria. Ele somente pode ser preenchido "
        "após a caracterização individual de cada contrato.",
    )
    legenda(doc, "Comparação entre os contratos analisados")
    tabela(
        doc,
        ["Dimensão de comparação", "DATAPREV", "SERPRO",
         "Convergência ou divergência"],
        [(d, PEND, PEND, "—") for d in [
            "Termo de Referência de origem", "Modelo de gestão contratual",
            "Estrutura de papéis", "Instâncias decisórias",
            "Mecanismos de coordenação", "Critérios de monitoramento",
            "Gestão de riscos",
        ]],
        proporcoes=[2.2, 1.2, 1.2, 1.8], fonte=9,
        alinhamentos=[None, "c", "c", "c"],
    )

    _quadro_eixo(
        doc, 1, "Papéis e responsabilidades",
        "O Eixo 1 verifica quem responde formalmente pelo contrato e se a "
        "estrutura de papéis cobre as funções exigidas. Interessa "
        "especialmente identificar a eventual participação de agentes nas "
        "decisões sem designação formal correspondente, situação que "
        "fragiliza a imputação de responsabilidade e compromete a "
        "rastreabilidade das decisões.",
        ["Gestor do contrato", "Fiscais designados",
         "Gestores das soluções, por item",
         "Participantes efetivos das decisões",
         "Atribuições definidas nos atos de designação",
         "Participantes das decisões sem formalização"],
    )
    _quadro_eixo(
        doc, 2, "Instâncias e níveis de decisão",
        "O Eixo 2 identifica os níveis de decisão e o tipo de situação que "
        "cada um resolve. O objetivo declarado no documento de orientação é "
        "compreender quem decide cada tipo de situação, o que pressupõe "
        "distinguir a competência formalmente atribuída da prática "
        "efetivamente observada nos autos.",
        ["Gestor do contrato", "Gestores das soluções", "Alta gestão",
         "Secretários", "Secretário-Executivo", "Ministro"],
    )
    _quadro_eixo(
        doc, 3, "Comunicação e coordenação",
        "O Eixo 3 examina como órgão e contratada se articulam, com que "
        "periodicidade, por quais canais e com que grau de formalização. "
        "Examina também a existência de registro das decisões, condição "
        "sem a qual a coordenação não produz trilha auditável.",
        ["Comunicação entre órgão e contratada", "Reuniões existentes",
         "Existência de comitês", "Registro das decisões",
         "Solicitação e aprovação de Ordens de Serviço",
         "Escalonamento de problemas e incidentes"],
    )
    _quadro_eixo(
        doc, 4, "Monitoramento e riscos",
        "O Eixo 4 examina como o contrato é acompanhado, quais indicadores "
        "existem, como são apurados e como os riscos são controlados. "
        "Neste produto, o exame recai sobre a existência e o funcionamento "
        "do mecanismo de monitoramento como elemento de governança; a "
        "proposição do modelo técnico de monitoramento orientado a riscos e "
        "resultados pertence ao Produto 03 e não é antecipada aqui.",
        ["Acompanhamento do contrato", "Indicadores existentes",
         "Controle de riscos", "Tratamento de problemas e incidentes",
         "Acompanhamento dos níveis de serviço",
         "Normas de cumprimento obrigatório"],
    )
    nota(
        doc, "Nota.",
        "A delimitação entre este produto e o Produto 03 quanto à matéria "
        "de monitoramento decorre de decisão registrada no Anexo C, sob o "
        "identificador D-04. O Produto 01 diagnostica o mecanismo "
        "existente; o Produto 03 propõe o modelo. A observância dessa "
        "delimitação evita duplicação de escopo e mantém a coerência entre "
        "os produtos.",
    )

    h(doc, "4.7. Classificação dos achados", 2)
    par(
        doc,
        "Os achados resultantes da aplicação dos critérios são classificados "
        "nas cinco categorias definidas pelo documento de orientação. As "
        "tabelas desta seção são os instrumentos de registro; cada linha "
        "corresponde a um achado, identificado por prefixo estável e "
        "vinculado à evidência que o sustenta.",
    )

    h(doc, "4.7.1. Lacunas — o que está faltando", 3)
    par(
        doc,
        "Registra-se como lacuna a ausência de elemento de governança "
        "exigido por norma, por instrumento contratual ou pela própria "
        "estrutura de gestão adotada pelo órgão.",
    )
    legenda(doc, "Lacunas identificadas")
    tabela(
        doc,
        ["ID", "Lacuna", "Contrato", "Evidência", "Severidade", "Situação"],
        [("L-01", "—", "—", "—", "—", PEND)],
        proporcoes=[0.6, 2.4, 0.9, 1.6, 1.0, 1.0], fonte=9,
        alinhamentos=["c", None, "c", "c", "c", "c"],
    )

    h(doc, "4.7.2. Sobreposições — quem faz a mesma coisa", 3)
    par(
        doc,
        "Registra-se como sobreposição a atribuição da mesma "
        "responsabilidade a mais de um papel, sem critério de distinção "
        "entre eles, com risco de duplicidade de esforço ou de diluição da "
        "responsabilidade.",
    )
    legenda(doc, "Sobreposições identificadas")
    tabela(
        doc,
        ["ID", "Sobreposição", "Contrato", "Papéis envolvidos",
         "Evidência", "Severidade", "Situação"],
        [("S-01", "—", "—", "—", "—", "—", PEND)],
        proporcoes=[0.6, 2.0, 0.9, 1.5, 1.3, 0.9, 0.9], fonte=8.5,
        alinhamentos=["c", None, "c", None, "c", "c", "c"],
    )

    h(doc, "4.7.3. Fragilidades — o que funciona mal", 3)
    par(
        doc,
        "Registra-se como fragilidade o elemento de governança que existe, "
        "porém se mostra insuficiente, inconsistente ou inefetivo para a "
        "finalidade a que se destina.",
    )
    legenda(doc, "Fragilidades identificadas")
    tabela(
        doc,
        ["ID", "Fragilidade", "Contrato", "Evidência", "Severidade",
         "Situação"],
        [("F-01", "—", "—", "—", "—", PEND)],
        proporcoes=[0.6, 2.4, 0.9, 1.6, 1.0, 1.0], fonte=9,
        alinhamentos=["c", None, "c", "c", "c", "c"],
    )

    h(doc, "4.7.4. Riscos — o que pode dar problema", 3)
    par(
        doc,
        "Registra-se como risco o evento futuro e incerto que, se "
        "concretizado, produz efeito negativo sobre a execução contratual, "
        "sobre a conformidade ou sobre o resultado pretendido. "
        "Probabilidade e impacto são sempre justificados, e o controle "
        "existente é avaliado quanto à efetividade.",
    )
    legenda(doc, "Riscos identificados")
    tabela(
        doc,
        ["ID", "Risco", "Causa raiz", "Probabilidade", "Impacto",
         "Controle existente", "Situação"],
        [("R-01", "—", "—", "—", "—", "—", PEND)],
        proporcoes=[0.6, 2.0, 1.6, 1.0, 0.9, 1.5, 0.9], fonte=8.5,
        alinhamentos=["c", None, None, "c", "c", "c", "c"],
    )

    h(doc, "4.7.5. Oportunidades — o que pode melhorar", 3)
    par(
        doc,
        "Registra-se como oportunidade a possibilidade de aprimoramento "
        "que não decorre de descumprimento nem de fragilidade, e cuja "
        "adoção eleva a eficiência ou a maturidade da gestão contratual.",
    )
    legenda(doc, "Oportunidades identificadas")
    tabela(
        doc,
        ["ID", "Oportunidade", "Contrato", "Ganho esperado", "Situação"],
        [("OP-01", "—", "—", "—", PEND)],
        proporcoes=[0.7, 2.6, 0.9, 2.0, 1.0], fonte=9,
        alinhamentos=["c", None, "c", "c", "c"],
    )

    h(doc, "4.8. Quadro-síntese dos achados", 2)
    par(
        doc,
        "O quadro-síntese consolida a contagem de achados por categoria e "
        "por severidade, permitindo a leitura imediata do resultado do "
        "diagnóstico. Será preenchido ao término da etapa 5 do "
        "procedimento descrito na seção 3.9.",
    )
    legenda(doc, "Síntese quantitativa dos achados")
    tabela(
        doc,
        ["Categoria", "Alta", "Média", "Baixa", "Total", "Situação"],
        [(c, "—", "—", "—", "—", PEND) for c in
         ["Lacunas", "Sobreposições", "Fragilidades", "Riscos",
          "Oportunidades", "Total geral"]],
        proporcoes=[2.0, 0.8, 0.8, 0.8, 0.8, 1.2], fonte=9,
        alinhamentos=[None, "c", "c", "c", "c", "c"],
    )
    nota(
        doc, "Nota.",
        "A ausência de achados registrados nesta versão não constitui "
        "resultado do diagnóstico e não deve ser interpretada como "
        "indicação de conformidade da governança dos contratos. Decorre "
        "unicamente da indisponibilidade da base documental, declarada na "
        "seção 4.1 e detalhada no Anexo A.",
    )


# --------------------------------------------------------------------------
# Capítulos 5 a 7 — Matrizes e recomendações
# --------------------------------------------------------------------------

def capitulo_5(doc):
    h(doc, "5. MATRIZ ANALÍTICA", 1, nova_pagina=True)
    par(
        doc,
        "A matriz analítica é entregável expressamente exigido pelo "
        "contrato de consultoria e pelo documento de orientação. Sua função "
        "é organizar, em uma única estrutura, os problemas identificados, "
        "as suas causas, as suas consequências e as melhorias possíveis, de "
        "modo que a passagem do diagnóstico para a recomendação seja "
        "verificável linha a linha.",
    )

    h(doc, "5.1. Estrutura e regra de preenchimento", 2)
    par(
        doc,
        "Cada linha da matriz corresponde a um problema e recebe "
        "identificador próprio, com o prefixo MA. O preenchimento observa "
        "quatro regras. Primeira: todo item da matriz tem origem em ao "
        "menos um achado do Capítulo 4, citado por identificador. Segunda: "
        "a causa registrada é a causa raiz, e não a manifestação do "
        "problema. Terceira: a consequência é descrita em termos de efeito "
        "verificável sobre a execução contratual, a conformidade ou o "
        "resultado, e não em termos genéricos. Quarta: a melhoria possível "
        "é enunciada de modo a permitir a sua conversão em recomendação "
        "formal no Capítulo 7, sem acréscimo de conteúdo novo.",
    )
    legenda(doc, "Campos da matriz analítica")
    tabela(
        doc,
        ["Campo", "Conteúdo", "Regra de preenchimento"],
        [
            ("ID", "Identificador do item, no formato MA-nn",
             "Sequencial e estável entre versões"),
            ("Problema", "Enunciado direto do problema identificado",
             "Deve corresponder a achado do Capítulo 4, citado por "
             "identificador"),
            ("Causa", "Causa raiz do problema",
             "Não confundir causa com sintoma nem com consequência"),
            ("Consequência", "Efeito verificável do problema",
             "Descrito em termos de execução, conformidade ou resultado"),
            ("Melhoria possível", "Ação de correção ou mitigação",
             "Enunciada de modo a permitir conversão direta em "
             "recomendação"),
            ("Evidência", "Documento que sustenta o problema",
             "Processo, peça e localização interna, quando aplicável"),
            ("Severidade", "Alta, média ou baixa",
             "Acompanhada da justificativa do enquadramento"),
        ],
        proporcoes=[1.1, 2.3, 2.7], fonte=9,
    )

    h(doc, "5.2. Matriz analítica", 2)
    par(
        doc,
        "A matriz será preenchida ao término da etapa 6 do procedimento "
        "descrito na seção 3.9, após a classificação dos achados.",
    )
    legenda(doc, "Matriz analítica")
    tabela(
        doc,
        ["ID", "Problema", "Causa", "Consequência", "Melhoria possível",
         "Evidência", "Severidade"],
        [("MA-01", "—", "—", "—", "—", "—", "—")],
        proporcoes=[0.7, 1.7, 1.4, 1.5, 1.7, 1.1, 0.9], fonte=8.5,
        alinhamentos=["c", None, None, None, None, "c", "c"],
    )


def capitulo_6(doc):
    h(doc, "6. MATRIZ DE ATRIBUIÇÃO DE RESPONSABILIDADES", 1,
      nova_pagina=True)
    par(
        doc,
        "A matriz de atribuição de responsabilidades, exigida pelo "
        "documento de orientação, é apresentada em duas versões. A primeira "
        "registra a situação apurada, isto é, o que está formalizado ou "
        "efetivamente praticado segundo evidência documental. A segunda "
        "registra a situação recomendada, que integra o conjunto de "
        "recomendações do produto.",
    )

    h(doc, "6.1. Convenção adotada", 2)
    par(
        doc,
        "A convenção é a definida no documento de orientação do Produto 01.",
    )
    legenda(doc, "Convenção de atribuição")
    tabela(
        doc,
        ["Letra", "Significado", "Observação"],
        [
            ("R", "Executa a atividade",
             "Pode haver mais de um responsável pela execução"),
            ("A", "Aprova a atividade e responde por ela",
             "Deve ser único por atividade; a duplicidade caracteriza "
             "sobreposição"),
            ("C", "É consultado antes da decisão",
             "Participação de caráter técnico ou consultivo"),
            ("I", "É informado após a decisão",
             "Não participa da deliberação"),
        ],
        proporcoes=[0.6, 2.4, 3.0], fonte=9,
        alinhamentos=["c", None, None],
    )
    nota(
        doc, "Nota.",
        "Na situação apurada, célula sem evidência documental é preenchida "
        "com a expressão “não identificado”, e nunca por presunção do que "
        "seria a atribuição natural do papel. A distinção entre o que está "
        "formalizado e o que é apenas praticado é registrada em coluna "
        "própria de observação, quando os autos permitirem essa "
        "diferenciação.",
    )

    atividades = [
        "Solicitar demanda e abrir Ordem de Serviço",
        "Aprovar Ordem de Serviço", "Executar o serviço",
        "Acompanhar a execução", "Apurar indicadores",
        "Emitir relatório de fiscalização", "Receber o objeto",
        "Atestar nota fiscal", "Tratar incidentes", "Escalar problemas",
        "Arquivar evidências", "Decidir sobre alterações contratuais",
    ]
    papeis = ["Gestor", "Fiscal técnico", "Fiscal administrativo",
              "Gestor da solução", "Alta gestão", "Preposto"]

    for i, nome in enumerate(["DATAPREV", "SERPRO"], start=2):
        h(doc, "6.%d. Situação apurada — contrato %s" % (i, nome), 2)
        par(
            doc,
            "O preenchimento depende dos atos de designação, da matriz de "
            "responsabilidades formalizada, se existente, e dos registros "
            "de execução contratual, todos relacionados no Anexo A.",
        )
        legenda(doc, "Atribuição de responsabilidades apurada: contrato %s"
                % nome)
        tabela(
            doc,
            ["Atividade do ciclo contratual"] + papeis,
            [(a,) + ("—",) * len(papeis) for a in atividades],
            proporcoes=[3.0] + [0.85] * len(papeis), fonte=8,
            alinhamentos=[None] + ["c"] * len(papeis),
        )

    h(doc, "6.4. Situação recomendada", 2)
    par(
        doc,
        "A matriz recomendada é elaborada após a conclusão do diagnóstico e "
        "integra as recomendações do Capítulo 7. Cada atribuição proposta "
        "declara expressamente o seu fundamento, com indicação da natureza, "
        "conforme a distinção entre exigência normativa e boa prática "
        "estabelecida na seção 1.2.",
    )
    legenda(doc, "Atribuição de responsabilidades recomendada")
    tabela(
        doc,
        ["Atividade do ciclo contratual", "Atribuição proposta",
         "Fundamento", "Natureza", "Situação"],
        [(a, "—", "—", "—", PEND) for a in atividades[:4]] +
        [("Demais atividades do ciclo", "—", "—", "—", PEND)],
        proporcoes=[2.4, 1.5, 1.5, 1.0, 1.0], fonte=9,
        alinhamentos=[None, "c", "c", "c", "c"],
    )


def capitulo_7(doc):
    h(doc, "7. RECOMENDAÇÕES TÉCNICAS", 1, nova_pagina=True)
    par(
        doc,
        "As recomendações técnicas constituem entregável exigido pelo "
        "contrato de consultoria, que as qualifica como estruturadas e "
        "mensuráveis. Este capítulo define o formato adotado, a regra de "
        "vinculação a achados e o quadro em que serão consolidadas.",
    )

    h(doc, "7.1. Formato e regra de vinculação", 2)
    par(
        doc,
        "O formato segue o definido no documento de orientação — problema, "
        "melhoria proposta, responsável e resultado esperado — acrescido de "
        "três campos exigidos pela metodologia desta consultoria: "
        "natureza, fundamento e indicador de implementação. O acréscimo "
        "responde à exigência contratual de que as recomendações sejam "
        "mensuráveis e à necessidade de distinguir obrigação de "
        "recomendação.",
    )
    legenda(doc, "Campos da recomendação")
    tabela(
        doc,
        ["Campo", "Conteúdo", "Regra"],
        [
            ("Problema", "Achado que a recomendação endereça",
             "Citado por identificador do Capítulo 4 ou da matriz "
             "analítica; recomendação sem problema correspondente é "
             "vedada"),
            ("Melhoria proposta", "Ação concreta a ser adotada",
             "Enunciada de modo verificável, sem formulação genérica"),
            ("Responsável", "Papel a quem cabe a ação",
             "Indicado por papel, nunca por nome de pessoa, salvo "
             "exigência do produto com base documental"),
            ("Resultado esperado", "Efeito pretendido",
             "Descrito em termos observáveis"),
            ("Natureza",
             "Exigência normativa ou boa prática",
             "Campo obrigatório; impede que boa prática seja apresentada "
             "como obrigação legal"),
            ("Fundamento", "Norma, se exigência; referencial, se boa "
                            "prática",
             "Citado somente se a identificação oficial estiver "
             "confirmada"),
            ("Indicador de implementação",
             "Forma de medir se a recomendação foi cumprida",
             "Deve permitir verificação objetiva por terceiro"),
            ("Prazo sugerido", "Imediato, curto ou médio prazo",
             "Acompanhado da justificativa do enquadramento"),
        ],
        proporcoes=[1.3, 2.0, 2.8], fonte=9,
    )
    nota(
        doc, "Ressalva.",
        "A regra de vinculação é a garantia de que o produto não "
        "recomendará práticas desvinculadas de problema real dos contratos "
        "examinados. Nenhuma recomendação será formulada a partir de "
        "referencial técnico sem achado correspondente identificado nos "
        "autos, conforme a vedação enunciada na seção 3.6.",
    )

    h(doc, "7.2. Quadro de recomendações", 2)
    par(
        doc,
        "O quadro será preenchido ao término da etapa 8 do procedimento "
        "descrito na seção 3.9.",
    )
    legenda(doc, "Recomendações técnicas")
    tabela(
        doc,
        ["ID", "Achado de origem", "Melhoria proposta", "Responsável",
         "Resultado esperado", "Natureza", "Indicador"],
        [("REC-01", "—", "—", "—", "—", "—", "—")],
        proporcoes=[0.8, 1.2, 1.9, 1.2, 1.7, 1.0, 1.2], fonte=8.5,
        alinhamentos=["c", "c", None, "c", None, "c", "c"],
    )


# --------------------------------------------------------------------------
# Capítulos 8 e 9
# --------------------------------------------------------------------------

def capitulo_8(doc):
    h(doc, "8. CONDIÇÕES PARA A CONCLUSÃO DO DIAGNÓSTICO", 1,
      nova_pagina=True)
    par(
        doc,
        "Este capítulo consolida, em um único lugar, tudo o que precisa "
        "ocorrer para que a Atividade 1.2 seja concluída. A consolidação "
        "tem finalidade prática: permitir que a supervisão e o órgão "
        "identifiquem, sem consulta a outras seções, exatamente quais "
        "providências destravam o trabalho e em que ordem.",
    )

    h(doc, "8.1. Dependências documentais", 2)
    par(
        doc,
        "As pendências abaixo condicionam a conclusão do diagnóstico. A "
        "relação nominal das peças correspondentes consta do Anexo A.",
    )
    legenda(doc, "Pendências documentais e efeito de cada uma")
    tabela(
        doc,
        ["ID", "Pendência", "Análise impedida", "Itens do Anexo A"],
        [
            ("P-01",
             "Atos de designação de gestor, fiscais e substitutos, nos dois "
             "contratos",
             "Eixo 1, na íntegra, e matriz de responsabilidades apurada",
             "S-05 a S-07; S-22 a S-24"),
            ("P-02",
             "Instrumentos contratuais completos, incluindo aditivos e "
             "termos de referência",
             "Caracterização do objeto e do modelo de gestão pactuado",
             "S-01 a S-04; S-18 a S-21"),
            ("P-03",
             "Registros das instâncias decisórias e das reuniões de "
             "acompanhamento", "Eixos 2 e 3",
             "S-08 a S-10; S-25 a S-27"),
            ("P-04",
             "Artefatos de fiscalização, recebimento e execução",
             "Eixo 4 e verificação do ciclo da demanda",
             "S-11 a S-13; S-28 a S-30"),
            ("P-05", "Mapas de gerenciamento de riscos dos contratos",
             "Avaliação do controle de riscos", "S-14; S-31"),
            ("P-06",
             "Definição da identificação oficial das referências "
             "normativas", "Fundamentação normativa dos achados",
             "Questão Q-01"),
            ("P-07",
             "Definição da versão dos referenciais de boas práticas",
             "Qualificação técnica dos achados e fundamentação das "
             "recomendações", "Questão Q-02"),
        ],
        proporcoes=[0.6, 2.3, 2.2, 1.4], fonte=9,
        alinhamentos=["c", None, None, "c"],
    )

    h(doc, "8.2. Questões abertas", 2)
    par(
        doc,
        "As questões abaixo dependem de definição da consultoria, da "
        "supervisão ou de ambas. O seu registro no produto atende à "
        "obrigação metodológica de tornar explícito o que ainda não está "
        "resolvido. O detalhamento consta do Anexo C.",
    )
    legenda(doc, "Questões abertas e efeito sobre o produto")
    tabela(
        doc,
        ["ID", "Questão", "Efeito sobre o produto", "Situação"],
        [
            ("Q-01",
             "Identificação oficial das referências normativas indicadas no "
             "documento de orientação",
             "Condiciona toda a fundamentação normativa do produto",
             "Aberta"),
            ("Q-02",
             "Versão dos referenciais COBIT e ITIL a adotar nos três "
             "produtos",
             "Condiciona a citação de práticas e a consistência entre os "
             "produtos", "Aberta"),
            ("Q-03",
             "Possibilidade de uso do processo indicado pela supervisão "
             "como referência metodológica",
             "Define se há fonte adicional legítima de referência",
             "Aberta"),
            ("Q-04",
             "Formato de entrega e requisitos de identidade visual do "
             "produto",
             "Define a etapa de produção do documento final",
             "Parcialmente resolvida"),
            ("Q-05",
             "Período de execução contratual que delimita a análise",
             "Define o universo documental a examinar e a "
             "representatividade da amostra", "Aberta"),
        ],
        proporcoes=[0.6, 2.4, 2.4, 1.1], fonte=9,
        alinhamentos=["c", None, None, "c"],
    )
    nota(
        doc, "Nota.",
        "A questão Q-04 encontra-se parcialmente resolvida quanto ao "
        "formato: o produto é entregue em arquivo de texto editável e em "
        "arquivo de leitura em formato portátil. Permanecem abertos os "
        "eventuais requisitos de identidade visual, tipografia, margens e "
        "numeração próprios da UNESCO ou do órgão, que, uma vez "
        "informados, serão aplicados sem alteração de conteúdo.",
    )

    h(doc, "8.3. Sequência de retomada", 2)
    par(
        doc,
        "Uma vez disponibilizados os documentos, a retomada observa a "
        "sequência abaixo, que corresponde às etapas 1 a 9 do procedimento "
        "descrito na seção 3.9 e não exige revisão da metodologia "
        "consolidada.",
    )
    legenda(doc, "Sequência de retomada do trabalho")
    tabela(
        doc,
        ["Ordem", "Etapa", "Pré-requisito"],
        [
            ("1", "Receber e analisar os instrumentos contratuais e os "
                  "termos de referência dos dois contratos",
             "Itens S-01, S-02, S-18 e S-19"),
            ("2", "Preencher a caracterização dos contratos, na seção 4.2",
             "Ordem 1 concluída"),
            ("3", "Receber e analisar os atos de designação",
             "Itens S-05 a S-07 e S-22 a S-24"),
            ("4", "Preencher o Eixo 1 e a matriz de responsabilidades "
                  "apurada", "Ordem 3 concluída"),
            ("5", "Consolidar a fundamentação normativa do produto",
             "Questões Q-01 e Q-02 resolvidas"),
            ("6", "Receber e analisar os registros de governança, "
                  "fiscalização e riscos",
             "Itens S-08 a S-16 e S-25 a S-33"),
            ("7", "Desenvolver o diagnóstico dos eixos 2, 3 e 4",
             "Ordens 4 e 6 concluídas"),
            ("8", "Consolidar a matriz analítica e as recomendações",
             "Ordem 7 concluída"),
            ("9", "Revisar tecnicamente e emitir a versão final do produto",
             "Ordem 8 concluída e lista de verificação aplicada"),
        ],
        proporcoes=[0.7, 3.2, 2.3], fonte=9,
        alinhamentos=["c", None, None],
    )
    par(
        doc,
        "A sequência preserva o prazo contratual de entrega, fixado em 12 "
        "de outubro de 2026, desde que as peças relacionadas nas duas "
        "primeiras faixas de prioridade do Anexo A sejam disponibilizadas "
        "com antecedência suficiente para a execução das ordens 1 a 4, que "
        "concentram o maior esforço de análise documental.",
    )


def capitulo_9(doc):
    h(doc, "9. CONCLUSÃO", 1, nova_pagina=True)
    blocos = [
        "O Produto 01 tem por finalidade diagnosticar a governança dos "
        "contratos estratégicos de TIC celebrados pelo Ministério da "
        "Educação com a DATAPREV e com o SERPRO, e cumpre duas atividades "
        "contratuais distintas. Esta versão entrega integralmente a "
        "primeira delas e deixa a segunda estruturada, porém não concluída, "
        "por razão externa à consultoria e integralmente documentada.",

        "A Atividade 1.1 está cumprida. O Capítulo 3 define o instrumento "
        "analítico completo, composto por duas dimensões de análise "
        "decompostas em onze eixos internos, vinte e seis critérios de "
        "avaliação com escalas próprias e regras de atribuição, nove "
        "categorias de fontes documentais, os parâmetros normativos "
        "aplicáveis com a respectiva situação de verificação, os "
        "referenciais de boas práticas com a regra que disciplina o seu "
        "uso, a escala de severidade, os padrões de registro por natureza "
        "de conteúdo, o sistema de identificadores que assegura a "
        "rastreabilidade e o procedimento de aplicação em nove etapas. O "
        "instrumento é aplicável de imediato e não depende de revisão.",

        "A Atividade 1.2 não pôde ser concluída. Nenhuma peça dos autos dos "
        "dois processos foi disponibilizada à consultoria até a data de "
        "elaboração deste documento. O Capítulo 4 registra essa situação "
        "por categoria documental e mantém montados todos os instrumentos "
        "de registro, de modo que o preenchimento possa ocorrer sem "
        "reestruturação. Os capítulos 5 a 7 preservam, pela mesma razão, a "
        "matriz analítica, as matrizes de atribuição de responsabilidades e "
        "o quadro de recomendações em estado de prontidão.",

        "Registra-se de forma expressa que a ausência de achados neste "
        "documento não constitui resultado do diagnóstico e não pode ser "
        "interpretada como indicação de conformidade, de adequação ou de "
        "maturidade da governança dos contratos examinados. Nada foi "
        "avaliado, porque nada foi disponibilizado para avaliação. A "
        "metodologia adotada veda a formulação de conclusão sem evidência, "
        "e essa vedação foi observada sem exceção, inclusive nos pontos em "
        "que a formulação de hipótese plausível seria tecnicamente "
        "possível.",

        "As condições para a conclusão estão consolidadas no Capítulo 8 e "
        "detalhadas no Anexo A. São de três ordens: a disponibilização das "
        "peças dos autos, na ordem de prioridade estabelecida; a "
        "confirmação, em fonte oficial, da identificação das referências "
        "normativas indicadas no documento de orientação; e a definição da "
        "versão dos referenciais de boas práticas e do período de execução "
        "contratual que delimita a análise. Nenhuma delas depende de "
        "decisão metodológica ainda por tomar.",

        "Atendidas essas condições, o diagnóstico poderá ser concluído "
        "dentro do prazo contratual, mediante a aplicação direta do "
        "instrumento ora entregue, sem alteração de escopo, de estrutura ou "
        "de critérios. Este documento é, nesse sentido, simultaneamente o "
        "cumprimento da primeira atividade contratual e o registro formal "
        "e auditável do que resta para o cumprimento da segunda.",
    ]
    for texto in blocos:
        par(doc, texto)


def assinatura(doc):
    h(doc, "ASSINATURA", 1, nova_pagina=True)
    par(
        doc,
        "Documento técnico elaborado no âmbito do Contrato de Consultoria "
        "UNESCO nº %s, em cumprimento às Atividades 1.1 e 1.2 previstas "
        "para o Produto 01." % CONTRATO,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run("____________________________________, "
                  "______ de ____________________ de 2026.")
    r.font.size = Pt(11)
    r.font.name = FONTE
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run("Local e data")
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = CINZA
    r2.font.name = FONTE

    for texto, tam, negrito, cor, antes in [
        ("_____________________________________________", 11, False,
         GRAFITE, 60),
        (CONSULTORA, 12, True, GRAFITE, 4),
        ("Consultora Individual", 10, False, CINZA, 0),
        ("Contrato de Consultoria UNESCO nº %s   ·   SA-1630/2026"
         % CONTRATO, 10, False, CINZA, 0),
    ]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(antes)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(texto)
        r.bold = negrito
        r.font.size = Pt(tam)
        r.font.color.rgb = cor
        r.font.name = FONTE


# --------------------------------------------------------------------------
# Anexos
# --------------------------------------------------------------------------

def anexo_a(doc):
    h(doc, "ANEXO A — DOCUMENTOS SOLICITADOS E PENDÊNCIAS", 1,
      nova_pagina=True)
    par(
        doc,
        "Este anexo relaciona nominalmente cada documento necessário à "
        "conclusão do diagnóstico, com o motivo da solicitação, a atividade "
        "a que se relaciona, o efeito concreto da sua ausência e a "
        "eventual alternativa de comprovação. Todos os itens encontram-se "
        "com situação aberta na data de elaboração deste documento. A "
        "identificação por código é estável e deve ser utilizada nas "
        "comunicações sobre o atendimento das solicitações.",
    )

    nova_secao(doc, paisagem=True)
    h(doc, "A.1. Peças dos autos dos processos analisados", 2)
    par(
        doc,
        "Os itens abaixo são solicitados em ambos os processos. A coluna "
        "de identificação registra o código correspondente em cada um "
        "deles.",
    )
    legenda(doc, "Peças solicitadas nos autos dos dois processos")
    autos = [
        ("S-01", "S-18",
         "Contrato celebrado, em versão assinada e completa, com todos os "
         "anexos",
         "Caracterizar o objeto, o modelo de gestão, as obrigações de "
         "fiscalização e as cláusulas de governança",
         "1.1 e 1.2 — todos os eixos",
         "Impede caracterizar o objeto e identificar as obrigações de "
         "gestão pactuadas; inviabiliza o diagnóstico", "Nenhuma"),
        ("S-02", "S-19", "Termo de Referência que fundamentou a contratação",
         "Identificar exigências de gestão, fiscalização, níveis de serviço "
         "e papéis previstos no planejamento", "1.1 e 1.2 — eixos 1, 3 e 4",
         "Impede comparar o previsto com o praticado",
         "Contrato, parcialmente"),
        ("S-03", "S-20", "Estudo Técnico Preliminar da contratação",
         "Compreender o contexto da decisão de contratar e o modelo de "
         "gestão concebido", "1.1 — contexto; 1.2 — eixo 2",
         "Limita a análise da dimensão estratégica",
         "Termo de Referência, parcialmente"),
        ("S-04", "S-21",
         "Termos aditivos, apostilamentos e demais alterações contratuais",
         "Verificar se as alterações impactaram a estrutura de gestão e de "
         "governança", "1.2 — eixos 1 e 2",
         "Pode gerar diagnóstico baseado em versão contratual superada",
         "Nenhuma"),
        ("S-05", "S-22",
         "Atos de designação do gestor do contrato e de seu substituto",
         "Identificar o papel formalizado, suas atribuições e a vigência da "
         "designação", "1.2 — eixo 1",
         "Impede a análise de papéis e a matriz de responsabilidades "
         "apurada", "Nenhuma"),
        ("S-06", "S-23",
         "Atos de designação dos fiscais técnico, administrativo e "
         "requisitante, e respectivos substitutos",
         "Identificar a cobertura dos papéis de fiscalização exigidos",
         "1.2 — eixo 1",
         "Impede avaliar se a estrutura de fiscalização está completa",
         "Nenhuma"),
        ("S-07", "S-24",
         "Documento que identifique os gestores das soluções por item "
         "contratado", "Mapear os responsáveis técnicos por solução",
         "1.2 — eixos 1 e 2",
         "Impede mapear a cadeia de decisão técnica",
         "Registros de Ordens de Serviço, parcialmente"),
        ("S-08", "S-25",
         "Atas, memórias ou registros das reuniões de acompanhamento, "
         "relativas a todo o período de execução",
         "Evidenciar a existência e o funcionamento das instâncias de "
         "coordenação", "1.2 — eixos 2 e 3",
         "Impede analisar comunicação, coordenação e registro de decisões",
         "Nenhuma"),
        ("S-09", "S-26",
         "Ato de instituição, regimento ou documento de funcionamento de "
         "comitê ou instância colegiada relacionada ao contrato, se houver",
         "Verificar a formalização da governança colegiada",
         "1.2 — eixos 2 e 3",
         "Não permite afirmar se a governança colegiada existe",
         "Atas, parcialmente"),
        ("S-10", "S-27",
         "Matriz de responsabilidades, fluxograma de processo ou documento "
         "equivalente que descreva o fluxo de gestão do contrato",
         "Comparar as responsabilidades formalizadas com as praticadas",
         "1.2 — eixo 1 e matriz de responsabilidades",
         "Impede a matriz de responsabilidades apurada com base documental",
         "Contrato e Termo de Referência, parcialmente"),
        ("S-11", "S-28",
         "Ordens de Serviço emitidas, em amostra representativa, incluídas "
         "as três mais recentes",
         "Verificar o fluxo real de solicitação, aprovação, execução e "
         "aceite", "1.2 — eixos 3 e 4",
         "Impede analisar o ciclo da demanda na prática", "Nenhuma"),
        ("S-12", "S-29",
         "Relatórios de fiscalização, de acompanhamento ou equivalentes "
         "emitidos no período",
         "Verificar como o contrato é efetivamente acompanhado",
         "1.2 — eixo 4", "Impede avaliar a efetividade da fiscalização",
         "Atestes, parcialmente"),
        ("S-13", "S-30",
         "Termos de recebimento provisório e definitivo emitidos, se "
         "houver", "Verificar o rito de recebimento praticado",
         "1.2 — eixo 4", "Impede avaliar o encerramento do ciclo de "
         "entrega", "Nenhuma"),
        ("S-14", "S-31",
         "Mapa de gerenciamento de riscos da contratação e suas "
         "atualizações",
         "Verificar a identificação e o tratamento formal de riscos",
         "1.2 — eixo 4", "Impede analisar o controle de riscos", "Nenhuma"),
        ("S-15", "S-32",
         "Registros de incidentes, problemas ou escalonamentos ocorridos na "
         "execução", "Verificar o tratamento real de problemas",
         "1.2 — eixos 3 e 4",
         "Impede analisar o mecanismo de escalonamento",
         "Atas, parcialmente"),
        ("S-16", "S-33",
         "Manifestações de órgãos de controle relativas ao contrato e as "
         "respectivas respostas do órgão, se houver",
         "Identificar apontamentos já existentes sobre a governança",
         "1.2 — todos os eixos",
         "Risco de o diagnóstico ignorar achado já constatado pelo "
         "controle", "Nenhuma"),
        ("S-17", "S-34",
         "Propostas comerciais ou propostas de atendimento apresentadas "
         "pela contratada",
         "Verificar o que foi pactuado em nível de execução, inclusive "
         "quanto ao acompanhamento", "1.2 — eixo 4",
         "Limita a análise do que foi efetivamente pactuado",
         "Contrato, parcialmente"),
    ]
    tabela(
        doc,
        ["DATAPREV", "SERPRO", "Documento solicitado", "Motivo",
         "Atividade relacionada", "Efeito da ausência", "Alternativa"],
        autos,
        proporcoes=[0.7, 0.7, 2.6, 2.6, 1.3, 2.6, 1.3], fonte=9,
        largura=25.5,
        alinhamentos=["c", "c", None, None, "c", None, "c"],
    )

    nova_secao(doc, paisagem=False)
    h(doc, "A.2. Documentos institucionais", 2)
    legenda(doc, "Documentos institucionais solicitados")
    tabela(
        doc,
        ["ID", "Documento solicitado", "Motivo", "Efeito da ausência"],
        [
            ("S-35",
             "Plano Diretor de Tecnologia da Informação e Comunicação "
             "2025–2027",
             "O contrato de consultoria exige alinhamento explícito a este "
             "plano",
             "Impede demonstrar o alinhamento institucional exigido "
             "contratualmente"),
            ("S-36",
             "Normativo interno do órgão sobre gestão e fiscalização de "
             "contratos de TIC, se houver",
             "Identificar as regras internas aplicáveis à governança "
             "contratual",
             "Pode gerar diagnóstico que ignora regra interna vigente"),
            ("S-37",
             "Organograma ou documento de estrutura da unidade de TIC do "
             "órgão",
             "Compreender a estrutura decisória institucional",
             "Limita o mapeamento dos níveis de decisão do Eixo 2"),
        ],
        proporcoes=[0.6, 2.3, 2.3, 2.4], fonte=9,
        alinhamentos=["c", None, None, None],
    )

    h(doc, "A.3. Ordem de prioridade para disponibilização", 2)
    par(
        doc,
        "A ordem abaixo foi estabelecida para destravar o trabalho com o "
        "menor número de interações sucessivas. As duas primeiras faixas "
        "concentram as peças de maior efeito sobre o andamento da análise.",
    )
    legenda(doc, "Prioridade de disponibilização")
    tabela(
        doc,
        ["Prioridade", "Itens", "Justificativa"],
        [
            ("1ª", "S-01, S-02, S-18 e S-19",
             "Sem os contratos e os termos de referência, nada pode ser "
             "caracterizado"),
            ("2ª", "S-05 a S-07; S-22 a S-24",
             "Destravam integralmente o Eixo 1 e a matriz de "
             "responsabilidades apurada"),
            ("3ª", "S-08 a S-10; S-25 a S-27",
             "Destravam os eixos 2 e 3"),
            ("4ª", "S-11 a S-15; S-28 a S-32", "Destravam o Eixo 4"),
            ("5ª", "S-03, S-04, S-16, S-17; S-20, S-21, S-33, S-34; "
                   "S-35 a S-37",
             "Complementam a análise e fornecem o contexto institucional"),
        ],
        proporcoes=[0.9, 2.2, 3.4], fonte=9,
        alinhamentos=["c", None, None],
    )

    h(doc, "A.4. Campos a completar na emissão", 2)
    par(
        doc,
        "Os campos abaixo integram a identificação do produto e não foram "
        "informados até a data de elaboração. Devem ser completados antes "
        "da emissão da versão assinada.",
    )
    legenda(doc, "Campos pendentes de informação")
    tabela(
        doc,
        ["Campo", "Onde consta", "Informação necessária"],
        [
            ("Unidade demandante", "Identificação do Produto",
             "Unidade do Ministério da Educação responsável pela demanda"),
            ("Projeto de cooperação técnica",
             "Capa e Identificação do Produto",
             "Identificação do projeto de cooperação a que a consultoria "
             "se vincula"),
            ("Local e data de emissão", "Bloco de assinatura",
             "Município e data da assinatura do produto"),
            ("Requisitos de identidade visual", "Formatação do documento",
             "Eventuais exigências de tipografia, margens, numeração ou "
             "elementos gráficos (Q-04)"),
        ],
        proporcoes=[1.5, 1.8, 3.2], fonte=9,
    )


def anexo_b(doc):
    h(doc, "ANEXO B — INSTRUMENTOS DE COLETA E REGISTRO", 1,
      nova_pagina=True)
    par(
        doc,
        "Este anexo reproduz os formulários que operacionalizam os padrões "
        "de registro definidos na seção 3.8. Seu preenchimento sistemático "
        "é a base da rastreabilidade de todo o produto: cada achado deve "
        "poder ser reconduzido, sem intermediação, ao documento que o "
        "sustenta.",
    )

    h(doc, "B.1. Ficha de documento analisado", 2)
    ficha(doc, [
        "FICHA DE DOCUMENTO ANALISADO",
        "",
        "Processo SEI .........................:",
        "Número do documento no SEI ...........:",
        "Data do documento ....................:",
        "Tipo de documento ....................:",
        "Contrato a que se refere .............:",
        "    ( ) DATAPREV   ( ) SERPRO   ( ) Ambos   ( ) Institucional",
        "",
        "Finalidade na análise ................:",
        "",
        "Informações relevantes encontradas:",
        "    1.",
        "    2.",
        "    3.",
        "",
        "Critério ou eixo a que se relaciona ..:",
        "Limitações da evidência ..............:",
        "Achados gerados (identificadores) ....:",
        "Data da análise ......................:",
    ])

    h(doc, "B.2. Ficha de achado", 2)
    ficha(doc, [
        "FICHA DE ACHADO",
        "",
        "Identificador ........................:",
        "Tipo .................................:",
        "    ( ) Lacuna         ( ) Sobreposição    ( ) Fragilidade",
        "    ( ) Risco          ( ) Oportunidade",
        "Contrato .............................:",
        "    ( ) DATAPREV       ( ) SERPRO          ( ) Ambos",
        "Eixo de análise ......................:",
        "Dimensão .............................:",
        "    ( ) Estratégica    ( ) Operacional",
        "",
        "Descrição do achado ..................:",
        "",
        "Evidência que o sustenta:",
        "    Documento ........................:",
        "    Processo SEI .....................:",
        "    Localização no documento .........:",
        "",
        "Fundamento normativo ou de boa prática:",
        "    Natureza .........................:",
        "        ( ) Exigência normativa    ( ) Boa prática",
        "    Referência .......................:",
        "",
        "Severidade ...........................:",
        "    ( ) Alta           ( ) Média            ( ) Baixa",
        "Justificativa da severidade ..........:",
        "Causa raiz ...........................:",
        "Consequência .........................:",
        "Recomendação vinculada ...............:",
        "",
        "Informação pendente que poderia alterar este achado:",
    ])

    h(doc, "B.3. Ficha de solicitação de documento", 2)
    ficha(doc, [
        "SOLICITAÇÃO DE DOCUMENTO",
        "",
        "Identificador ........................:",
        "Documento solicitado .................:",
        "Processo SEI onde deve ser localizado :",
        "Motivo da solicitação ................:",
        "Atividade relacionada ................:",
        "Efeito da ausência ...................:",
        "Alternativas de comprovação ..........:",
        "Prioridade ...........................:",
        "    ( ) 1ª    ( ) 2ª    ( ) 3ª    ( ) 4ª    ( ) 5ª",
        "Data da solicitação ..................:",
        "Situação .............................:",
        "    ( ) Aberta                     ( ) Atendida",
        "    ( ) Não localizado nos autos   ( ) Prejudicada",
    ])

    h(doc, "B.4. Registro de documentos analisados", 2)
    par(
        doc,
        "O registro consolidado abaixo é preenchido a cada documento "
        "recebido e constitui a base de rastreabilidade do produto. "
        "Nenhum documento foi analisado até a data de elaboração deste "
        "texto.",
    )
    legenda(doc, "Registro de documentos analisados")
    tabela(
        doc,
        ["#", "Documento", "Processo SEI", "Nº no SEI", "Data",
         "Finalidade na análise", "Limitações da evidência"],
        [("—", "—", "—", "—", "—", "—", "—")],
        proporcoes=[0.5, 2.0, 1.4, 1.0, 0.8, 2.0, 1.8], fonte=8.5,
        alinhamentos=["c", None, "c", "c", "c", None, None],
    )


def anexo_c(doc):
    h(doc, "ANEXO C — QUESTÕES ABERTAS E DECISÕES REGISTRADAS", 1,
      nova_pagina=True)
    par(
        doc,
        "Este anexo registra as questões ainda não resolvidas e as decisões "
        "metodológicas já tomadas no curso da consultoria. O registro das "
        "decisões tem função prática: preserva a coerência entre os três "
        "produtos e impede que opções já firmadas sejam reabertas sem "
        "justificativa.",
    )

    nova_secao(doc, paisagem=True)
    h(doc, "C.1. Questões abertas", 2)
    legenda(doc, "Questões abertas")
    tabela(
        doc,
        ["ID", "Questão", "Por que importa", "Quem decide", "Situação"],
        [
            ("Q-01",
             "O documento de orientação cita uma portaria sem indicação de "
             "ano, uma instrução normativa com ano possivelmente "
             "divergente e um acórdão sem identificação do colegiado "
             "prolator. Qual é a identificação oficial correta de cada "
             "uma?",
             "A fundamentação normativa de todo o produto depende disso; "
             "citar norma com identificação incorreta compromete o "
             "documento",
             "Consultora, com consulta à supervisão e verificação em fonte "
             "oficial", "Aberta"),
            ("Q-02",
             "Qual versão dos referenciais será adotada e mantida nos três "
             "produtos, quanto ao COBIT e quanto ao ITIL?",
             "Consistência entre os produtos e precisão na citação de "
             "práticas e objetivos", "Consultora", "Aberta"),
            ("Q-03",
             "O processo indicado pela supervisão é referido como de "
             "consulta apenas. Seu material pode servir como referência "
             "metodológica, sem figurar como evidência, ou está "
             "integralmente vedado?",
             "Define se há fonte adicional legítima de referência "
             "metodológica",
             "Consultora, com validação da supervisão", "Aberta"),
            ("Q-04",
             "O produto deve ser entregue apenas em formato de leitura ou "
             "também em formato editável? Há requisito de identidade "
             "visual, tipografia, margens ou numeração específico?",
             "Afeta a etapa de produção do documento final",
             "Consultora, com consulta à supervisão",
             "Parcialmente resolvida quanto ao formato; identidade visual "
             "em aberto"),
            ("Q-05",
             "Há período de execução contratual de referência para a "
             "análise, com data de corte definida?",
             "Define o universo documental a examinar e evita análise "
             "incompleta ou desatualizada", "Consultora", "Aberta"),
        ],
        proporcoes=[0.6, 3.4, 2.4, 1.8, 1.6], fonte=9, largura=25.5,
        alinhamentos=["c", None, None, None, "c"],
    )

    nova_secao(doc, paisagem=False)
    h(doc, "C.2. Decisões registradas", 2)
    legenda(doc, "Decisões metodológicas registradas")
    tabela(
        doc,
        ["ID", "Decisão", "Justificativa"],
        [
            ("D-01",
             "Escopo restrito aos contratos com DATAPREV e SERPRO, nos três "
             "produtos",
             "Confirmação expressa da consultora; o contrato de consultoria "
             "não nomeia contratos específicos, e o documento de orientação "
             "nomeia apenas estes dois"),
            ("D-02",
             "Modelo de referência utilizado exclusivamente como padrão de "
             "estrutura, formatação e apresentação",
             "O modelo é produto de outra consultoria, com objeto diverso; "
             "seu conteúdo não constitui evidência desta análise"),
            ("D-03",
             "Adoção de documento base único para o controle do projeto",
             "Reduz o risco de versões desencontradas e concentra o "
             "controle em um só ponto"),
            ("D-04",
             "Delimitação entre o Produto 01 e o Produto 03 quanto ao "
             "monitoramento: aquele diagnostica o mecanismo existente, este "
             "propõe o modelo",
             "Evita duplicação e sobreposição de escopo entre produtos"),
            ("D-05",
             "Adoção de marcação de situação das informações, com "
             "indicação de fonte documental para o que já está consolidado",
             "Garante rastreabilidade e impede que hipótese se converta em "
             "fato ao longo das etapas"),
            ("D-06",
             "Emprego da expressão “não localizado nos autos” em lugar de "
             "“inexistente”, salvo quando houver base para afirmar a "
             "inexistência",
             "Rigor metodológico: a ausência de acesso não comprova a "
             "ausência do artefato"),
            ("D-07",
             "Reclassificação como pendentes dos dados que, em versão "
             "anterior do documento base, haviam sido extraídos do modelo "
             "de referência",
             "Conformidade com a decisão D-02"),
        ],
        proporcoes=[0.6, 2.8, 3.2], fonte=9,
        alinhamentos=["c", None, None],
    )


def anexo_d(doc):
    h(doc, "ANEXO D — REFERÊNCIAS", 1, nova_pagina=True)
    par(
        doc,
        "As referências abaixo estão organizadas em três blocos: os "
        "documentos da própria consultoria, o conjunto normativo e de "
        "orientação de controle, e os referenciais técnicos. A situação de "
        "verificação de cada item é reproduzida para que o leitor "
        "identifique, de imediato, o que já está apto a citação e o que "
        "depende de confirmação.",
    )

    h(doc, "D.1. Documentos da consultoria", 2)
    tabela(
        doc,
        ["Documento", "Natureza", "Uso neste produto"],
        [
            ("Contrato de prestação de serviços UNESCO nº %s e termos de "
             "referência anexos" % CONTRATO, "Instrumento contratual",
             "Fonte das exigências, atividades, produtos e prazos"),
            ("Documento de orientação do Produto 01",
             "Direcionamento da supervisão",
             "Fonte dos eixos de análise e do formato das matrizes"),
            ("Documento base do projeto de consultoria, versão 1.1",
             "Instrumento interno de controle",
             "Consolidação metodológica e controle de pendências"),
        ],
        proporcoes=[2.6, 1.6, 2.6], fonte=9,
    )

    h(doc, "D.2. Normas e orientações de controle", 2)
    tabela(
        doc,
        ["Referência", "Situação de verificação", "Condição de citação"],
        [
            ("Constituição da República Federativa do Brasil, artigo 37",
             "Identificação confirmada", "Citação autorizada"),
            ("Lei nº 14.133/2021", "Identificação confirmada",
             "Citação autorizada; indicação de dispositivos na etapa de "
             "análise documental"),
            ("Lei nº 4.320/1964", "Identificação confirmada",
             "Citação autorizada, restrita ao instituto da liquidação da "
             "despesa"),
            ("Súmula TCU nº 269", "Identificação confirmada",
             "Citação autorizada no âmbito do Eixo 4"),
            ("Norma do SISP sobre contratação de soluções de TIC", VALID,
             "Citação vedada até a confirmação oficial (Q-01)"),
            ("Portaria referida no documento de orientação, sem indicação "
             "de ano", VALID,
             "Citação vedada até a confirmação oficial (Q-01)"),
            ("Instrução normativa referida no documento de orientação, com "
             "ano divergente", VALID,
             "Citação vedada até a confirmação oficial (Q-01)"),
            ("Acórdão referido no documento de orientação, sem "
             "identificação do colegiado", VALID,
             "Citação vedada até a confirmação oficial (Q-01)"),
            ("Plano Diretor de TIC 2025–2027", PEND,
             "Documento não disponibilizado (item S-35)"),
        ],
        proporcoes=[3.0, 1.5, 2.5], fonte=9,
    )

    h(doc, "D.3. Referenciais técnicos", 2)
    tabela(
        doc,
        ["Referencial", "Escopo de uso neste produto", "Situação"],
        [
            ("COBIT",
             "Governança e gestão de TIC: responsabilidades, estrutura "
             "decisória, riscos, controles e monitoramento",
             "Versão a definir (Q-02)"),
            ("ITIL",
             "Gestão de serviços de TIC: ciclo de demanda, incidentes, "
             "problemas, mudanças, níveis de serviço e melhoria contínua",
             "Versão a definir (Q-02)"),
        ],
        proporcoes=[1.2, 3.8, 1.5], fonte=9,
    )
    nota(
        doc, "Nota.",
        "Os dados bibliográficos completos — número, data, ementa, veículo "
        "de publicação e versão dos referenciais — serão incorporados a "
        "esta relação após a verificação em fonte oficial, na forma da "
        "regra enunciada na seção 3.5. Até que essa verificação ocorra, a "
        "referência permanece identificada apenas no nível em que pode ser "
        "afirmada com segurança.",
    )


# --------------------------------------------------------------------------
# Montagem
# --------------------------------------------------------------------------

def limpar_paragrafos_orfaos(doc):
    """Remove o parágrafo vazio que separa uma tabela do título seguinte
    quando esse título já força quebra de página. Sem essa limpeza, o
    parágrafo vazio sobra sozinho e produz uma página em branco."""
    corpo = doc.element.body
    removidos = 0
    filhos = list(corpo)
    for atual, seguinte in zip(filhos, filhos[1:]):
        if atual.tag != qn("w:p") or seguinte.tag != qn("w:p"):
            continue
        if "".join(atual.itertext()).strip():
            continue
        pPr = atual.find(qn("w:pPr"))
        if pPr is not None and pPr.find(qn("w:sectPr")) is not None:
            continue  # carrega quebra de seção: preservar
        pPr_seguinte = seguinte.find(qn("w:pPr"))
        if pPr_seguinte is None:
            continue
        quebra = pPr_seguinte.find(qn("w:pageBreakBefore"))
        if quebra is None:
            continue
        if quebra.get(qn("w:val")) in (None, "true", "1", "on"):
            corpo.remove(atual)
            removidos += 1
    return removidos


def propriedades(doc):
    nucleo = doc.core_properties
    nucleo.title = "%s — %s" % (PRODUTO, TITULO)
    nucleo.subject = ("Consultoria Técnica em Governança de Contratos de "
                      "TIC — UNESCO · Ministério da Educação")
    nucleo.author = CONSULTORA
    nucleo.category = "Documento técnico de consultoria"
    nucleo.comments = (
        "Produto 01 — Atividades contratuais 1.1 e 1.2. Contratos "
        "analisados: DATAPREV (Processo SEI nº %s) e SERPRO (Processo SEI "
        "nº %s)." % (SEI_DATAPREV, SEI_SERPRO)
    )
    nucleo.keywords = ("governança contratual; contratos de TIC; DATAPREV; "
                       "SERPRO; diagnóstico")


def main():
    _CONTADOR_QUADRO[0] = 0
    doc = Document()
    configurar_estilos(doc)
    secao = configurar_pagina(doc)
    montar_cabecalho(secao)
    montar_rodape(secao)
    propriedades(doc)

    capa(doc)
    identificacao(doc)
    sumario(doc)
    siglas(doc)
    sumario_executivo(doc)
    capitulo_1(doc)
    capitulo_2(doc)
    capitulo_3(doc)
    capitulo_4(doc)
    capitulo_5(doc)
    capitulo_6(doc)
    capitulo_7(doc)
    capitulo_8(doc)
    capitulo_9(doc)
    assinatura(doc)
    anexo_a(doc)
    anexo_b(doc)
    anexo_c(doc)
    anexo_d(doc)

    vazios = limpar_paragrafos_orfaos(doc)
    if vazios:
        print("Parágrafos órfãos removidos: %d" % vazios)

    doc.save(SAIDA_DOCX)
    print("DOCX gerado: %s" % SAIDA_DOCX)


if __name__ == "__main__":
    main()
