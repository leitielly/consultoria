#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Confere se alguma coluna de tabela do .docx é estreita demais para a
maior palavra indivisível que nela aparece.

Usa as mesmas métricas do gerador, de modo que as duas medições não podem
divergir. Executado após cada geração, como verificação de qualidade.
"""

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Emu

from metricas_fonte import MARGEM_CELULA_CM, largura_cm, maior_palavra


def margem_da_tabela(tabela):
    """Recuo interno efetivo, somados os dois lados, em centímetros."""
    margens = tabela._tbl.tblPr.find(qn("w:tblCellMar"))
    if margens is None:
        return MARGEM_CELULA_CM
    total = 0.0
    for lado in ("left", "right"):
        elem = margens.find(qn("w:" + lado))
        total += int(elem.get(qn("w:w"))) if elem is not None else 108
    return total / 1440.0 * 2.54


def verificar(caminho):
    doc = Document(caminho)
    problemas = []
    for n, tabela in enumerate(doc.tables, start=1):
        margem = margem_da_tabela(tabela)
        for i, coluna in enumerate(tabela.columns):
            larg_col = coluna.cells[0].width
            if not larg_col:
                continue
            util = Emu(larg_col).cm - margem
            for j, celula in enumerate(coluna.cells):
                for p in celula.paragraphs:
                    for run in p.runs:
                        palavra = maior_palavra(run.text)
                        if not palavra:
                            continue
                        tam = run.font.size.pt if run.font.size else 9
                        larg = largura_cm(palavra, tam, bool(run.font.bold))
                        if larg > util + 0.005:
                            problemas.append((n, i, j, palavra, larg, util))
    return problemas


def main():
    caminho = sys.argv[1] if len(sys.argv) > 1 else (
        "Produto_01_Diagnostico_Governanca_Contratual.docx")
    achados = verificar(caminho)
    if not achados:
        print("Colunas: nenhuma palavra seria partida.")
        return 0
    vistos = set()
    for tab, col, lin, palavra, larg, util in achados:
        chave = (tab, col, palavra)
        if chave in vistos:
            continue
        vistos.add(chave)
        print("tabela %-3d coluna %-2d linha %-3d  %-22s precisa %.2f cm, "
              "tem %.2f cm" % (tab, col, lin, palavra, larg, util))
    print("\n%d ocorrências distintas." % len(vistos))
    return 1


if __name__ == "__main__":
    sys.exit(main())
